"""Shared ProfitTrack brand + sheet builders for the reseller toolkit workbooks.

Everything here is written to survive Excel, Google Sheets, the Excel/Sheets
phone apps and Apple Numbers: numeric helper columns instead of TEXT() or
wildcard criteria, simple conditional-format rule types only, and one-section
number formats on anything a chart reads.
"""
from openpyxl.chart import BarChart, LineChart
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.comments import Comment
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.drawing.text import Font as DFont
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------- Brand ----------
NAVY = "0B1F3A"
NAVY_2 = "16325C"
TEAL = "14B8A6"
TEAL_DARK = "0F766E"
TEAL_LIGHT = "CCFBF1"
INPUT_FILL = "F0FDFA"   # very light teal = cells the user types in
CALC_FILL = "F1F5F9"    # light slate = formula cells (don't type here)
YELLOW = "FEF9C3"       # settings the user may change
WHITE = "FFFFFF"
SLATE = "475569"
RED = "DC2626"
RED_LIGHT = "FEE2E2"
AMBER = "B45309"
AMBER_LIGHT = "FEF3C7"
GREEN = "15803D"
GREEN_LIGHT = "DCFCE7"
INPUT_BLUE = "1E3A8A"
FONT = "Arial"

CUR = '$#,##0.00;[Red]-$#,##0.00;"-"'
CUR_SIMPLE = "$#,##0.00"     # for cells a chart reads (phone previews ignore 3-section formats)
CUR0 = '$#,##0;[Red]-$#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE_FMT = "yyyy-mm-dd"
MONTHS = ["Jan / Ene", "Feb / Feb", "Mar / Mar", "Apr / Abr", "May / May", "Jun / Jun",
          "Jul / Jul", "Aug / Ago", "Sep / Sep", "Oct / Oct", "Nov / Nov", "Dec / Dic"]
PLATFORMS = ["eBay", "Facebook Marketplace", "Etsy", "Poshmark", "Other"]
YES, NO = "Yes / Sí", "No"

thin = Side(style="thin", color="CBD5E1")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def paint(ws, rng, color):
    for row in ws[rng]:
        for c in row:
            c.fill = fill(color)


def note(text, author="ProfitTrack"):
    return Comment(text, author)


def banner(ws, last_col, title, subtitle):
    """Navy title bar in rows 1-3 (+ teal rule in row 4), used on every sheet except the cover."""
    paint(ws, f"A1:{last_col}3", NAVY)
    ws.merge_cells(f"B1:{last_col}1")
    ws["B1"] = "ProfitTrack"
    ws["B1"].font = font(9, True, TEAL)
    ws["B1"].alignment = Alignment(vertical="bottom")
    ws.merge_cells(f"B2:{last_col}2")
    ws["B2"] = title
    ws["B2"].font = font(18, True, WHITE)
    ws["B2"].alignment = Alignment(vertical="center")
    ws.merge_cells(f"B3:{last_col}3")
    ws["B3"] = subtitle
    ws["B3"].font = font(10, False, "C7D2FE", italic=True)
    ws["B3"].alignment = Alignment(vertical="top")
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 20
    paint(ws, f"A4:{last_col}4", TEAL)
    ws.row_dimensions[4].height = 4
    ws.column_dimensions["A"].width = 2.5
    ws.sheet_view.showGridLines = False


def header_cell(c, text):
    c.value = text
    c.font = font(10, True, WHITE)
    c.fill = fill(NAVY_2)
    c.alignment = CENTER
    c.border = BORDER


def section_title(ws, cell_rng, text):
    ws.merge_cells(cell_rng)
    c = ws[cell_rng.split(":")[0]]
    c.value = text
    c.font = font(12, True, NAVY)
    c.alignment = Alignment(vertical="center")
    c.border = Border(bottom=Side(style="medium", color=TEAL))


def setting_cell(c, value, fmt=None):
    """Yellow, unlocked cell the user may change."""
    c.value = value
    c.font = font(11, True, INPUT_BLUE)
    c.fill = fill(YELLOW)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.protection = Protection(locked=False)
    if fmt:
        c.number_format = fmt


def contains_rule(ws, rng_, text, fnt, fll=None):
    """'Cell contains text' highlight -- a simple rule type Apple Numbers keeps on import."""
    first = rng_.split(":")[0]
    rule = Rule(type="containsText", operator="containsText", text=text,
                dxf=DifferentialStyle(font=fnt, fill=fll))
    rule.formula = [f'NOT(ISERROR(SEARCH("{text}",{first})))']
    ws.conditional_formatting.add(rng_, rule)


def list_validation(ws, rng_, items=None, source=None, title="List / Lista"):
    formula = source if source else '"' + ",".join(items) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, errorTitle=title,
                        error="Choose an option from the list.\nElige una opción de la lista.")
    ws.add_data_validation(dv)
    dv.add(rng_)
    return dv


def number_validation(ws, ranges, kind="decimal", low="0", high=None, title="Number / Número",
                      error="Enter a number ≥ 0.\nEscribe un número ≥ 0."):
    if high is None:
        dv = DataValidation(type=kind, operator="greaterThanOrEqual", formula1=low, allow_blank=True,
                            errorTitle=title, error=error)
    else:
        dv = DataValidation(type=kind, operator="between", formula1=low, formula2=high, allow_blank=True,
                            errorTitle=title, error=error)
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)
    return dv


def date_validation(ws, rng_):
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        errorTitle="Date / Fecha",
                        error="Enter a valid date (e.g. 2026-03-15).\nEscribe una fecha válida (ej. 2026-03-15).")
    ws.add_data_validation(dv)
    dv.add(rng_)
    return dv


def style_grid(ws, headers, header_row, first, last, money=(), pct=(), dates=(), centered=(), ints=()):
    """Header row + formatted empty grid. headers: [(col, text, width, 'in'|'calc')]."""
    ws.row_dimensions[header_row].height = 48
    for col, text, width, _ in headers:
        header_cell(ws[f"{col}{header_row}"], text)
        ws.column_dimensions[col].width = width
    for r in range(first, last + 1):
        for col, _, _, kind in headers:
            c = ws[f"{col}{r}"]
            c.fill = fill(INPUT_FILL if kind == "in" else CALC_FILL)
            c.border = BORDER
            c.font = font(10, color=INPUT_BLUE if kind == "in" else "000000")
            if col in money:
                c.number_format = CUR
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col in pct:
                c.number_format = "0%"
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col in dates:
                c.number_format = DATE_FMT
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col in ints:
                c.number_format = '#,##0.0;[Red]-#,##0.0;"-"'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col in centered:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")


def hide_helpers(ws, header_row, cols_labels, first, last):
    for col, label in cols_labels:
        ws[f"{col}{header_row}"] = f"Helper – do not edit: {label}"
        ws[f"{col}{header_row}"].font = font(8, False, "94A3B8")
        ws.column_dimensions[col].hidden = True
        for r in range(first, last + 1):
            ws[f"{col}{r}"].font = font(8, color="94A3B8")


def kpi_cards(ws, top, cards):
    """cards: [(c1, c2, EN, ES, formula, fmt, sub_formula, sub_fmt)] -> 4-row navy tiles starting at `top`."""
    ws.row_dimensions[top].height = 20
    ws.row_dimensions[top + 1].height = 16
    ws.row_dimensions[top + 2].height = 38
    ws.row_dimensions[top + 3].height = 18
    for c1, c2, en, es, f, fmt, sub, sub_fmt in cards:
        paint(ws, f"{c1}{top}:{c2}{top + 2}", NAVY)
        paint(ws, f"{c1}{top + 3}:{c2}{top + 3}", TEAL)
        if c1 != c2:
            for k in range(4):
                ws.merge_cells(f"{c1}{top + k}:{c2}{top + k}")
        ws[f"{c1}{top}"] = en
        ws[f"{c1}{top}"].font = font(9, True, TEAL)
        ws[f"{c1}{top + 1}"] = es
        ws[f"{c1}{top + 1}"].font = font(8, False, "C7D2FE", italic=True)
        ws[f"{c1}{top + 2}"] = f
        ws[f"{c1}{top + 2}"].font = font(22, True, WHITE)
        ws[f"{c1}{top + 2}"].number_format = fmt
        ws[f"{c1}{top + 3}"] = sub
        ws[f"{c1}{top + 3}"].font = font(7, False, WHITE)
        if sub_fmt:
            ws[f"{c1}{top + 3}"].number_format = sub_fmt
        for k in range(4):
            ws[f"{c1}{top + k}"].alignment = Alignment(horizontal="center", vertical="center")


def table_rows(ws, first, rows, cols, fmts, bold_first=True):
    """Write a striped summary table. rows: list of lists of values/formulas aligned with cols."""
    for i, values in enumerate(rows):
        r = first + i
        for j, (col, v) in enumerate(zip(cols, values)):
            c = ws[f"{col}{r}"]
            c.value = v
            c.border = BORDER
            c.font = font(10, bold=(bold_first and j == 0))
            c.fill = fill(TEAL_LIGHT if i % 2 == 0 else WHITE)
            if fmts.get(col):
                c.number_format = fmts[col]
            if fmts.get(col) in ("0", "#,##0"):
                c.alignment = Alignment(horizontal="center")


def total_row(ws, r, cols, values, fmts):
    for col, v in zip(cols, values):
        c = ws[f"{col}{r}"]
        c.value = v
        c.font = font(10, True, WHITE)
        c.fill = fill(NAVY)
        c.border = BORDER
        if fmts.get(col):
            c.number_format = fmts[col]
        if fmts.get(col) in ("0", "#,##0"):
            c.alignment = Alignment(horizontal="center")


def small_title(text):
    cp = CharacterProperties(sz=1200, b=True, solidFill=NAVY, latin=DFont(typeface=FONT))
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def axis_text(size=800):
    cp = CharacterProperties(sz=size)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def bar_chart(title, horizontal=False, y_title=None, labels=True, money=True):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = small_title(title)
    ch.style = 10
    ch.legend = None
    if y_title:
        ch.y_axis.title = y_title
    ch.y_axis.numFmt = "$#,##0" if money else "0"
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showSerName = False
        ch.dataLabels.showCatName = False
        ch.dataLabels.showLegendKey = False
    return ch


def color_series(series, color, line=None):
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = line or color


def build_cover(wb, line_en, line_es, contents, compat, kicker="RESELLER TOOLKIT  •  KIT PARA REVENDEDORES",
                platforms="eBay  •  Facebook Marketplace  •  Etsy  •  Poshmark  •  Mercari  •  Depop"):
    cv = wb.active
    cv.title = "Cover"
    cv.sheet_view.showGridLines = False
    for col in range(1, 13):
        cv.column_dimensions[get_column_letter(col)].width = 11
    paint(cv, "A1:L48", NAVY)
    paint(cv, "A13:L13", TEAL)
    cv.row_dimensions[13].height = 5
    cv.merge_cells("B4:K4")
    cv["B4"] = kicker
    cv["B4"].font = font(11, True, TEAL)
    cv["B4"].alignment = CENTER
    cv.merge_cells("B6:K8")
    cv["B6"] = "ProfitTrack"
    cv["B6"].font = font(48, True, WHITE)
    cv["B6"].alignment = CENTER
    cv.merge_cells("B9:K9")
    cv["B9"] = line_en
    cv["B9"].font = font(18, True, TEAL_LIGHT)
    cv["B9"].alignment = CENTER
    cv.merge_cells("B10:K10")
    cv["B10"] = line_es
    cv["B10"].font = font(14, False, "C7D2FE", italic=True)
    cv["B10"].alignment = CENTER
    cv.merge_cells("B11:K11")
    cv["B11"] = platforms
    cv["B11"].font = font(10, False, "94A3B8")
    cv["B11"].alignment = CENTER
    for r in (6, 7, 8):
        cv.row_dimensions[r].height = 22
    cv.row_dimensions[9].height = 28
    cv.row_dimensions[10].height = 22
    cv.merge_cells("B15:K15")
    cv["B15"] = "WHAT'S INSIDE  /  CONTENIDO   (click to open • clic para abrir)"
    cv["B15"].font = font(12, True, TEAL)
    cv["B15"].alignment = CENTER
    r = 17
    for sheet, label, desc in contents:
        cv.merge_cells(f"C{r}:J{r}")
        c = cv[f"C{r}"]
        c.value = label
        c.hyperlink = f"#'{sheet}'!A1"
        c.font = Font(name=FONT, size=13, bold=True, color=WHITE, underline="single")
        c.alignment = LEFT
        cv.merge_cells(f"C{r + 1}:J{r + 1}")
        d = cv[f"C{r + 1}"]
        d.value = desc
        d.font = font(9, False, "94A3B8")
        d.alignment = LEFT
        cv.row_dimensions[r].height = 22
        cv.row_dimensions[r + 1].height = 26
        r += 3
    r += 1
    cv.merge_cells(f"B{r}:K{r}")
    cv[f"B{r}"] = compat
    cv[f"B{r}"].font = font(9, False, "94A3B8")
    cv[f"B{r}"].alignment = CENTER
    r += 2
    cv.merge_cells(f"B{r}:K{r}")
    cv[f"B{r}"] = "Start here → Quick Start Guide   |   Empieza aquí → Guía Rápida"
    cv[f"B{r}"].font = Font(name=FONT, size=11, bold=True, color=TEAL, underline="single")
    cv[f"B{r}"].hyperlink = "#'Quick Start Guide'!A1"
    cv[f"B{r}"].alignment = CENTER
    r += 3
    cv.merge_cells(f"B{r}:K{r}")
    cv[f"B{r}"] = ("© ProfitTrack  •  v1.0  •  For personal use only – do not resell or redistribute / "
                   "Solo para uso personal – prohibida su reventa o distribución")
    cv[f"B{r}"].font = font(8, False, "64748B")
    cv[f"B{r}"].alignment = CENTER
    return cv


DEVICE_ROWS = [
    ("💻 Computer (Windows / Mac): open in Microsoft Excel, or in Google Sheets from any browser. Everything works.",
     "💻 Computadora (Windows / Mac): ábrelo en Microsoft Excel o en Google Sheets desde cualquier navegador. Todo funciona."),
    ("📱 iPhone, iPad & Android — RECOMMENDED: the free Microsoft Excel app or the free Google Sheets app. Formulas, dropdowns, charts and alerts all work.",
     "📱 iPhone, iPad y Android — RECOMENDADO: la app gratuita Microsoft Excel o la app gratuita Google Sheets. Fórmulas, listas, gráficas y alertas funcionan."),
    ("🍎 Apple Numbers: formulas, Dashboard, charts and alerts work, but Numbers removes dropdown lists from Excel files (you'll see a message – this is normal). Type the option exactly as listed, or re-add a menu: select the cells → Format → Cell → Data Format → Pop-Up Menu.",
     "🍎 Apple Numbers: las fórmulas, el Panel, las gráficas y las alertas funcionan, pero Numbers elimina las listas desplegables de los archivos de Excel (verás un aviso – es normal). Escribe la opción tal como aparece, o vuelve a crear el menú: selecciona las celdas → Formato → Celda → Formato de datos → Menú desplegable."),
]


def build_guide(wb, steps, sections):
    """steps: [(en, es)]; sections: [(title, [(en, es)] )] rendered after the numbered steps."""
    qs = wb.create_sheet("Quick Start Guide")
    banner(qs, "E", "Quick Start Guide  /  Guía Rápida", "Up and running in 5 minutes.  •  Listo en 5 minutos.")
    qs.column_dimensions["B"].width = 6
    qs.column_dimensions["C"].width = 58
    qs.column_dimensions["D"].width = 58
    qs.column_dimensions["E"].width = 2
    qs.row_dimensions[6].height = 22
    header_cell(qs["B6"], "#")
    header_cell(qs["C6"], "ENGLISH")
    header_cell(qs["D6"], "ESPAÑOL")

    def row_height(en, es):
        return max(30, 14 * (max(len(en), len(es)) // 62 + 1) + 8)

    r = 7
    for i, (en, es) in enumerate(steps, start=1):
        qs[f"B{r}"] = i
        qs[f"B{r}"].font = font(14, True, WHITE)
        qs[f"B{r}"].fill = fill(TEAL_DARK)
        qs[f"B{r}"].alignment = CENTER
        qs[f"C{r}"], qs[f"D{r}"] = en, es
        for col in "CD":
            qs[f"{col}{r}"].font = font(10)
            qs[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            qs[f"{col}{r}"].fill = fill(TEAL_LIGHT if i % 2 else WHITE)
        for col in "BCD":
            qs[f"{col}{r}"].border = BORDER
        qs.row_dimensions[r].height = row_height(en, es)
        r += 1

    legend = ("Color Legend  /  Leyenda de Colores", [
        ("__INPUT__", "Light teal = type your data here", "Verde claro = escribe tus datos aquí"),
        ("__CALC__", "Grey = automatic formula – don't type here", "Gris = fórmula automática – no escribas aquí"),
        ("__YELLOW__", "Yellow = settings you can change", "Amarillo = ajustes que puedes cambiar"),
    ])
    all_sections = [legend] + list(sections) + [("Phone, Tablet & Computer  /  Celular, Tableta y Computadora", DEVICE_ROWS)]
    colors = {"__INPUT__": INPUT_FILL, "__CALC__": CALC_FILL, "__YELLOW__": YELLOW}
    for title, rows in all_sections:
        r += 1
        qs.merge_cells(f"B{r}:D{r}")
        qs[f"B{r}"] = title
        qs[f"B{r}"].font = font(12, True, NAVY)
        qs[f"B{r}"].border = Border(bottom=Side(style="medium", color=TEAL))
        r += 1
        for row in rows:
            if len(row) == 3:
                key, en, es = row
                qs[f"B{r}"].fill = fill(colors[key])
            else:
                en, es = row
            qs[f"C{r}"], qs[f"D{r}"] = en, es
            for col in "BCD":
                qs[f"{col}{r}"].border = BORDER
            for col in "CD":
                qs[f"{col}{r}"].font = font(10)
                qs[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            qs.row_dimensions[r].height = row_height(en, es) if len(row) == 2 else 18
            r += 1
    return qs


def finish(wb, dashboard, out, title):
    if dashboard is not None:
        dashboard.protection.sheet = True
        dashboard.protection.selectLockedCells = False
        dashboard.protection.selectUnlockedCells = False
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("Cover", "Quick Start Guide", "Settings", "Materials") else TEAL
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0
    wb.properties.title = title
    wb.properties.creator = "ProfitTrack"
    wb.save(out)
    print("saved", out)
