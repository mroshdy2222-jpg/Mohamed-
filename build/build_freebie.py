"""Build the free lead magnet: ProfitTrack Flip Calculator (bilingual EN/ES)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from common import (AMBER, AMBER_LIGHT, BORDER, CUR, FONT, GREEN, GREEN_LIGHT, NAVY, PCT, RED, RED_LIGHT, SLATE, TEAL,
                    TEAL_DARK, TEAL_LIGHT, WHITE, banner, build_cover, contains_rule, fill, finish, font, header_cell,
                    list_validation, note, number_validation, section_title, setting_cell, style_grid)

OUT = "product/FREE_ProfitTrack_Flip_Calculator.xlsx"
FIRST, LAST = 6, 30
SHOP_URL = "https://YOURSHOP.etsy.com"   # replace with your Etsy shop link before uploading

wb = Workbook()
build_cover(
    wb, "Free Flip Profit Calculator", "Calculadora Gratis de Ganancia por Reventa",
    [
        ("Flip Calculator", "🧮  Flip Calculator / Calculadora", "Is it worth buying? Know your profit before you pay / ¿Vale la pena? Sabe tu ganancia antes de pagar"),
        ("Fees", "⚙️  Fees / Comisiones", "Typical platform fees – edit to match yours / Comisiones típicas – edítalas"),
        ("More Tools", "🧰  More Tools / Más Herramientas", "The full ProfitTrack reseller toolkit / El kit completo"),
    ],
    "FREE • GRATIS  —  💻 PC / Mac  •  📱 iPhone / Android  —  Excel • Google Sheets",
)

# ---------- Fees ----------
fe = wb.create_sheet("Fees")
banner(fe, "E", "Platform Fees  /  Comisiones por Plataforma",
       "Typical fees (US, 2026). Fees change – edit the yellow cells to match what you pay.  •  "
       "Comisiones típicas (EE. UU., 2026). Edita las celdas amarillas con lo que tú pagas.")
for col, w in {"B": 26, "C": 13, "D": 13, "E": 60}.items():
    fe.column_dimensions[col].width = w
fe.row_dimensions[6].height = 30
for col, text in zip("BCDE", ["Platform\nPlataforma", "Fee %\nComisión %", "Fixed Fee $\nTarifa Fija $", "Notes\nNotas"]):
    header_cell(fe[f"{col}6"], text)
fees = [
    ("eBay", 0.1325, 0.40, "Most categories; fixed fee per order. / La mayoría de categorías."),
    ("Poshmark", 0.20, 0, "Sales $15+. Under $15: flat $2.95 → set 0% and $2.95. / Ventas de $15+."),
    ("Etsy", 0.095, 0.45, "6.5% transaction + 3% processing; $0.25 + $0.20 listing. / Transacción + procesamiento."),
    ("Facebook Marketplace", 0.10, 0, "Shipped orders. Local cash pickup: use 'Local / Cash'. / Pedidos con envío."),
    ("Mercari", 0.10, 0.50, "Check Mercari's current fees. / Revisa las comisiones actuales de Mercari."),
    ("Local / Cash", 0, 0, "No fees on local cash sales. / Sin comisión en ventas locales."),
    ("Other / Otro", 0.10, 0, "Enter your rate. / Escribe tu comisión."),
]
for i, (p, pct, fix, n) in enumerate(fees):
    r = 7 + i
    fe[f"B{r}"] = p
    fe[f"B{r}"].font = font(10, True, NAVY)
    fe[f"B{r}"].border = BORDER
    setting_cell(fe[f"C{r}"], pct, "0.00%")
    setting_cell(fe[f"D{r}"], fix, "$#,##0.00")
    fe[f"E{r}"] = n
    fe[f"E{r}"].font = font(9, False, SLATE, italic=True)
    fe[f"E{r}"].border = BORDER
PLATS = [f[0] for f in fees]
section_title(fe, "B16:D16", "Your Rule  /  Tu Regla")
fe["B17"] = "Good flip = margin at least / Buena reventa = margen mínimo"
fe["B17"].font = font(10, True, NAVY)
fe.merge_cells("B17:C17")
setting_cell(fe["D17"], 0.30, "0%")
GOOD = "Fees!$D$17"
number_validation(fe, ["C7:C13", "D17"], high="0.9", title="Percent", error="0%–90%")
number_validation(fe, ["D7:D13"])

# ---------- Flip Calculator ----------
fc = wb.create_sheet("Flip Calculator", 1)
banner(fc, "N", "Flip Calculator  /  Calculadora de Reventa",
       "Type what you'd pay and what it sells for. Compare up to 25 items before you buy.  •  "
       "Escribe cuánto pagarías y en cuánto se vende. Compara hasta 25 artículos antes de comprar.")
headers = [
    ("B", "Item\nArtículo", 28, "in"),
    ("C", "Platform\nPlataforma", 20, "in"),
    ("D", "Buy Price\nPrecio de Compra", 13, "in"),
    ("E", "Sale Price\nPrecio de Venta", 13, "in"),
    ("F", "Shipping You Pay\nEnvío que Pagas", 14, "in"),
    ("G", "Other Costs\nOtros Costos", 12, "in"),
    ("H", "Fee %\nComisión %", 10, "calc"),
    ("I", "Fees $\nComisiones $", 12, "calc"),
    ("J", "Profit\nGanancia", 12, "calc"),
    ("K", "Margin\nMargen", 10, "calc"),
    ("L", "ROI\nRetorno", 10, "calc"),
    ("M", "Verdict\nVeredicto", 26, "calc"),
]
style_grid(fc, headers, 5, FIRST, LAST, money="DEFGIJ", pct="HKL", centered="M")
FP, FPCT, FFIX = "Fees!$B$7:$B$13", "Fees!$C$7:$C$13", "Fees!$D$7:$D$13"
for r in range(FIRST, LAST + 1):
    fc[f"H{r}"] = f'=IF(OR(B{r}="",E{r}=""),"",IFERROR(INDEX({FPCT},MATCH(C{r},{FP},0)),0))'
    fc[f"H{r}"].number_format = "0.0%"
    fc[f"I{r}"] = f'=IF(H{r}="","",E{r}*H{r}+IFERROR(INDEX({FFIX},MATCH(C{r},{FP},0)),0))'
    fc[f"J{r}"] = f'=IF(I{r}="","",E{r}-D{r}-F{r}-G{r}-I{r})'
    fc[f"K{r}"] = f'=IF(J{r}="","",IF(E{r}=0,0,J{r}/E{r}))'
    fc[f"K{r}"].number_format = PCT
    fc[f"L{r}"] = f'=IF(J{r}="","",IF(D{r}+G{r}=0,"∞",J{r}/(D{r}+G{r})))'
    fc[f"L{r}"].number_format = '0%;[Red]-0%;"-"'
    fc[f"M{r}"] = (f'=IF(J{r}="","",IF(J{r}<0,"⚠ LOSS / PÉRDIDA",IF(K{r}>={GOOD},"✔ GOOD FLIP / BUENA REVENTA",'
                   f'"● THIN MARGIN / MARGEN BAJO")))')
samples = [("Nike Air Max 90", "eBay", 22, 89.99, 12.45, 0),
           ("Levi's 501 Jeans", "Poshmark", 6, 45, 0, 0),
           ("Pyrex Bowl Set", "Local / Cash", 8, 40, 0, 0),
           ("Funko Pop", "Mercari", 14, 22, 4.50, 0),
           ("Bose Headphones", "eBay", 95, 140, 10, 0)]
for i, s in enumerate(samples):
    for col, v in zip("BCDEFG", s):
        fc[f"{col}{FIRST + i}"] = v
fc["B6"].comment = note("Rows 6–10 are EXAMPLES – type over them.\nLas filas 6–10 son EJEMPLOS – escribe encima.")
list_validation(fc, f"C{FIRST}:C{LAST}", source="=Fees!$B$7:$B$13", title="Platform / Plataforma")
number_validation(fc, [f"D{FIRST}:G{LAST}"])
contains_rule(fc, f"M{FIRST}:M{LAST}", "GOOD", Font(name=FONT, bold=True, color=GREEN), fill(GREEN_LIGHT))
contains_rule(fc, f"M{FIRST}:M{LAST}", "THIN", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(fc, f"M{FIRST}:M{LAST}", "LOSS", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
fc.freeze_panes = "C6"
r = LAST + 2
fc.merge_cells(f"B{r}:M{r}")
fc[f"B{r}"] = ("Want automatic dashboards, inventory alerts and tax-ready expense & mileage logs? → See the 'More Tools' tab.  /  "
               "¿Quieres paneles automáticos, alertas de inventario y registro de gastos y millaje? → Mira la pestaña 'More Tools'.")
fc[f"B{r}"].font = font(10, True, TEAL_DARK)
fc[f"B{r}"].hyperlink = "#'More Tools'!A1"

# ---------- More Tools ----------
mt = wb.create_sheet("More Tools")
banner(mt, "D", "The Full ProfitTrack Toolkit  /  El Kit Completo",
       "Liked this calculator? These tools do the rest – same design, same English + Español layout.  •  "
       "¿Te gustó? Estas herramientas hacen el resto.")
for col, w in {"B": 34, "C": 70, "D": 3}.items():
    mt.column_dimensions[col].width = w
tools = [
    ("📊 Profit Tracker", "Log every sale – fees, net profit & margin automatic, low-margin alerts, dashboard by platform & month. / Registro de ventas con panel automático."),
    ("📦 Inventory Tracker", "1,000 items, days held, Fresh → Watch → Stale alerts, stock value, sell-through, 10 oldest items. / Inventario con alertas."),
    ("🧾 Expenses & Mileage", "Reseller expense categories, receipt check, mileage at the IRS rate by trip date, total tax deductions. / Gastos y millaje."),
    ("🧰 Reseller Bundle", "All three tools together for less. / Las tres herramientas juntas por menos."),
]
r = 6
for name, desc in tools:
    mt[f"B{r}"] = name
    mt[f"B{r}"].font = font(13, True, WHITE)
    mt[f"B{r}"].fill = fill(NAVY)
    mt[f"B{r}"].alignment = Alignment(vertical="center", indent=1)
    mt[f"C{r}"] = desc
    mt[f"C{r}"].font = font(10)
    mt[f"C{r}"].fill = fill(TEAL_LIGHT)
    mt[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    mt.row_dimensions[r].height = 44
    r += 1
r += 1
mt.merge_cells(f"B{r}:C{r}")
mt[f"B{r}"] = f"👉  Get them here / Consíguelas aquí:  {SHOP_URL}"
mt[f"B{r}"].hyperlink = SHOP_URL
mt[f"B{r}"].font = Font(name=FONT, size=14, bold=True, color=TEAL_DARK, underline="single")
mt.row_dimensions[r].height = 30
mt[f"B{r + 2}"] = "Free for personal use. Please share the download link, not the file. / Gratis para uso personal. Comparte el enlace, no el archivo."
mt[f"B{r + 2}"].font = font(8, False, SLATE, italic=True)

finish(wb, None, OUT, "ProfitTrack Flip Calculator (Free)")
