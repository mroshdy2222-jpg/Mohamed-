"""Convert the listing Markdown into a simple, clean .docx."""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import sys

SRC = sys.argv[1] if len(sys.argv) > 1 else "product/ProfitTrack_Listing_Copy.md"
OUT = SRC[:-3] + ".docx"
NAVY, TEAL = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x0F, 0x76, 0x6E)
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Arial", Pt(10.5)
ARABIC = re.compile(r"[؀-ۿ]")


def add_runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part:
            p.add_run(part)
    if ARABIC.search(text):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        i += 1
        buf = []
        while not lines[i].startswith("```"):
            buf.append(lines[i]); i += 1
        t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
        cell = t.rows[0].cells[0]; cell.text = ""
        p = cell.paragraphs[0]
        for j, b in enumerate(buf):
            r = p.add_run(b); r.font.size = Pt(9.5)
            if j < len(buf) - 1:
                r.add_break()
        doc.add_paragraph()
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[\s\-|]+\|$", lines[i]):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i); cell.text = ""
                add_runs(cell.paragraphs[0], val)
        doc.add_paragraph()
        continue
    elif ln.startswith("#"):
        level = len(ln) - len(ln.lstrip("#"))
        h = doc.add_heading(ln.lstrip("#").strip(), level=min(level, 3))
        for r in h.runs:
            r.font.color.rgb = NAVY if level == 1 else TEAL
            r.font.name = "Arial"
    elif ln.startswith("> "):
        p = doc.add_paragraph(); add_runs(p, ln[2:])
        for r in p.runs:
            r.italic = True
    elif ln.startswith("- [ ] "):
        p = doc.add_paragraph("☐ "); add_runs(p, ln[6:])
    elif ln.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); add_runs(p, ln[2:])
    elif ln.strip() == "---" or not ln.strip():
        pass
    else:
        p = doc.add_paragraph(); add_runs(p, ln)
    i += 1
doc.save(OUT)
print("saved", OUT)
