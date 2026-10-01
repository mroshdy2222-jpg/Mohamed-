"""Build ProfitTrack Inventory — Reseller Inventory Tracker (bilingual EN/ES)."""
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import Reference
from openpyxl.styles import Alignment, Font

from common import (AMBER, AMBER_LIGHT, BORDER, CUR, CUR0, CUR_SIMPLE, DATE_FMT, FONT, GREEN, GREEN_LIGHT, NAVY,
                    PCT, PLATFORMS, RED, RED_LIGHT, SLATE, TEAL, TEAL_DARK, TEAL_LIGHT, WHITE, axis_text, banner,
                    bar_chart, build_cover, build_guide, color_series, contains_rule, date_validation, fill, finish,
                    font, header_cell, hide_helpers, kpi_cards, list_validation, note, number_validation,
                    section_title, setting_cell, style_grid, table_rows, total_row)

OUT = "product/ProfitTrack_Inventory_Tracker.xlsx"
FIRST, LAST = 6, 1005          # 1,000 items
CAT_FIRST, CAT_LAST = 8, 19    # Settings: 10 categories + 2 custom

IN_STOCK, LISTED, SOLD = "In Stock / En Inventario", "Listed / Publicado", "Sold / Vendido"
STATUSES = [IN_STOCK, LISTED, SOLD, "Donated / Donado", "Damaged / Dañado"]
CATEGORIES = [
    "Clothing / Ropa",
    "Shoes / Zapatos",
    "Bags & Accessories / Bolsos y Accesorios",
    "Jewelry & Watches / Joyería y Relojes",
    "Home & Kitchen / Hogar y Cocina",
    "Electronics / Electrónicos",
    "Toys & Collectibles / Juguetes y Coleccionables",
    "Books & Media / Libros y Medios",
    "Vintage & Antiques / Vintage y Antigüedades",
    "Other / Otro",
]

wb = Workbook()
build_cover(
    wb, "Reseller Inventory Tracker", "Control de Inventario para Revendedores",
    [
        ("Dashboard", "📊  Dashboard / Panel", "Stock value, stale items, sell-through, oldest items / Valor del inventario, artículos estancados, tasa de venta"),
        ("Inventory", "📦  Inventory / Inventario", "Every item from purchase to sale – days held & alerts automatic / Cada artículo de la compra a la venta – días y alertas automáticos"),
        ("Settings", "⚙️  Settings / Configuración", "Stale-item alert days & your categories / Días de alerta y tus categorías"),
        ("Quick Start Guide", "🚀  Quick Start Guide / Guía Rápida", "Step-by-step setup in 5 minutes / Configuración paso a paso en 5 minutos"),
    ],
    "💻 PC / Mac  •  📱 iPhone / Android  —  Excel • Google Sheets   |   Computadora y celular — Excel • Google Sheets",
)

# =====================================================================
# SETTINGS
# =====================================================================
st = wb.create_sheet("Settings")
banner(st, "F", "Settings  /  Configuración", "Alert thresholds and item categories.  •  Límites de alerta y categorías.")
for col, w in {"B": 40, "C": 14, "D": 3, "E": 3, "F": 50}.items():
    st.column_dimensions[col].width = w
section_title(st, "B6:C6", "Age Alerts  /  Alertas de Antigüedad")
WATCH, STALE = "Settings!$C$7", "Settings!$C$8"
for r, label, val in ((7, "'Watch' after (days) / 'Vigilar' después de (días)", 60),
                      (8, "'Stale' after (days) / 'Estancado' después de (días)", 90)):
    st[f"B{r}"] = label
    st[f"B{r}"].font = font(10, True, NAVY)
    st[f"B{r}"].border = BORDER
    setting_cell(st[f"C{r}"], val, "0")
st.merge_cells("B9:C11")
st["B9"] = ("Unsold items older than these limits are flagged on the Inventory sheet and counted on the Dashboard.\n"
            "Los artículos sin vender más antiguos que estos límites se marcan en Inventario y se cuentan en el Panel.")
st["B9"].font = font(8, False, SLATE, italic=True)
st["B9"].alignment = Alignment(wrap_text=True, vertical="top")
number_validation(st, ["C7:C8"], kind="whole", low="1", high="3650", title="Days / Días",
                  error="Enter a number of days (1–3650).\nEscribe un número de días (1–3650).")

section_title(st, "F6:F6", "Item Categories  /  Categorías")
st.row_dimensions[7].height = 36
header_cell(st["F7"], "Category (edit or add your own)\nCategoría (edita o agrega)")
for i in range(CAT_LAST - CAT_FIRST + 1):
    c = st[f"F{CAT_FIRST + i}"]
    c.value = CATEGORIES[i] if i < len(CATEGORIES) else None
    c.border = BORDER
    c.fill = fill("FEF9C3")
    c.font = font(10, color="1E3A8A")
st[f"F{CAT_LAST + 1}"] = "Rows 18–19 are free for your own categories. / Las filas 18–19 son para tus propias categorías."
st[f"F{CAT_LAST + 1}"].font = font(8, False, SLATE, italic=True)

# =====================================================================
# INVENTORY
# =====================================================================
iv = wb.create_sheet("Inventory", 1)
banner(iv, "Q", "Inventory  /  Inventario",
       "One row per item. Type in the light-teal columns only – grey columns calculate automatically.  •  "
       "Una fila por artículo. Escribe solo en las columnas verde claro – las grises se calculan solas.")
headers = [
    ("B", "SKU\nCódigo", 11, "in"),
    ("C", "Item Name\nArtículo", 28, "in"),
    ("D", "Category\nCategoría", 26, "in"),
    ("E", "Source\nFuente", 16, "in"),
    ("F", "Purchase Date\nFecha de Compra", 13, "in"),
    ("G", "Cost\nCosto", 11, "in"),
    ("H", "Platform\nPlataforma", 18, "in"),
    ("I", "List Price\nPrecio Publicado", 12, "in"),
    ("J", "Status\nEstado", 21, "in"),
    ("K", "Sold Date\nFecha de Venta", 13, "in"),
    ("L", "Sold Price\nPrecio de Venta", 12, "in"),
    ("M", "Location / Bin\nUbicación", 12, "in"),
    ("N", "Days Held\nDías", 9, "calc"),
    ("O", "Potential Profit\nGanancia Potencial", 13, "calc"),
    ("P", "Gross Profit\nGanancia Bruta", 12, "calc"),
    ("Q", "Age Alert\nAlerta", 30, "calc"),
]
style_grid(iv, headers, 5, FIRST, LAST, money="GILOP", dates="FK", centered="BJMNQ")
for r in range(FIRST, LAST + 1):
    iv[f"S{r}"] = f'=IF(C{r}="",0,IF(OR(J{r}="",J{r}="{IN_STOCK}",J{r}="{LISTED}"),1,0))'
    iv[f"T{r}"] = f'=IF(C{r}="",0,IF(J{r}="{SOLD}",1,0))'
    iv[f"U{r}"] = f'=IF(S{r}=1,IF(ISNUMBER(F{r}),TODAY()-F{r},-1),-1)'
    iv[f"V{r}"] = f'=IF(U{r}<0,IF(S{r}=1,5,0),IF(U{r}<=30,1,IF(U{r}<=60,2,IF(U{r}<=90,3,4))))'
    iv[f"W{r}"] = f'=IF(U{r}<0,-1E+15,U{r}+ROW()/1000000)'
    iv[f"X{r}"] = f'=IF(T{r}=1,IF(AND(ISNUMBER(F{r}),ISNUMBER(K{r})),K{r}-F{r},0),0)'
    iv[f"Y{r}"] = f'=IF(T{r}=1,IF(AND(ISNUMBER(F{r}),ISNUMBER(K{r})),1,0),0)'
    iv[f"Z{r}"] = f'=IF(U{r}>={STALE},1,0)'
    iv[f"AA{r}"] = f'=IF(S{r}=1,IF(J{r}="{LISTED}",0,1),0)'
    iv[f"AB{r}"] = f'=IF(T{r}=1,IF(ISNUMBER(L{r}),L{r}-G{r},0),0)'
    iv[f"AC{r}"] = f'=IF(S{r}=1,IF(ISNUMBER(I{r}),I{r}-G{r},0),0)'
    iv[f"N{r}"] = f'=IF(S{r}=1,IF(U{r}<0,"",U{r}),IF(Y{r}=1,X{r},""))'
    iv[f"N{r}"].number_format = "0"
    iv[f"O{r}"] = f'=IF(S{r}=1,IF(ISNUMBER(I{r}),I{r}-G{r},""),"")'
    iv[f"P{r}"] = f'=IF(T{r}=1,IF(ISNUMBER(L{r}),L{r}-G{r},""),"")'
    iv[f"Q{r}"] = (f'=IF(C{r}="","",IF(T{r}=1,"✔ SOLD / VENDIDO",IF(S{r}=0,"— CLOSED / CERRADO",'
                   f'IF(U{r}<0,"⚠ ADD DATE / AGREGA FECHA",IF(U{r}>={STALE},"⚠ STALE – REPRICE / ESTANCADO",'
                   f'IF(U{r}>={WATCH},"● WATCH / VIGILAR","✔ FRESH / RECIENTE"))))))')
hide_helpers(iv, 5, [("S", "unsold"), ("T", "sold"), ("U", "age"), ("V", "age bucket"), ("W", "oldest key"),
                     ("X", "days to sell"), ("Y", "has days to sell"), ("Z", "stale"), ("AA", "not listed"),
                     ("AB", "gross profit"), ("AC", "potential profit")], FIRST, LAST)
iv["N5"].comment = note("Unsold: days since purchase (today − purchase date). Sold: days from purchase to sale.\n"
                        "Sin vender: días desde la compra. Vendido: días de la compra a la venta.")
iv["O5"].comment = note("Unsold items: List Price − Cost (before fees & shipping).\nSin vender: Precio Publicado − Costo (antes de comisiones y envío).")
iv["P5"].comment = note("Sold items: Sold Price − Cost (before fees & shipping). For net profit use ProfitTrack Profit Tracker.\n"
                        "Vendidos: Precio de Venta − Costo (antes de comisiones y envío).")
iv["J5"].comment = note("Blank, In Stock or Listed = still in inventory. Sold, Donated or Damaged = closed.\n"
                        "Vacío, En Inventario o Publicado = sigue en inventario. Vendido, Donado o Dañado = cerrado.")

samples = [  # sku, item, cat, source, purchased, cost, platform, list, status, sold date, sold price, bin
    ("A-001", "Levi's 501 Jeans 32x30", 0, "Goodwill", date(2026, 3, 8), 6.00, 3, 38.00, 1, None, None, "Bin 1"),
    ("A-002", "Nike Air Max 90 – Size 10", 1, "Garage sale", date(2026, 3, 15), 15.00, 0, 95.00, 2, date(2026, 4, 2), 89.99, "Shelf A"),
    ("A-003", "Coach Leather Crossbody", 2, "Estate sale", date(2026, 4, 1), 20.00, 3, 120.00, 1, None, None, "Shelf B"),
    ("A-004", "Pyrex Butterprint Bowl", 4, "Thrift store", date(2026, 4, 12), 8.00, 1, 45.00, 2, date(2026, 5, 3), 40.00, "Shelf C"),
    ("A-005", "Vintage Polaroid SX-70", 5, "Flea market", date(2026, 5, 2), 45.00, 0, 220.00, 2, date(2026, 5, 20), 210.00, "Shelf D"),
    ("A-006", "Lululemon Align Leggings", 0, "Goodwill", date(2026, 5, 10), 10.00, 3, 48.00, 1, None, None, "Bin 1"),
    ("A-007", "LEGO Star Wars Lot", 6, "Facebook", date(2026, 5, 18), 60.00, 0, 150.00, 1, None, None, "Shelf E"),
    ("A-008", "Harry Potter Hardcover Set", 7, "Library sale", date(2026, 6, 1), 12.00, 0, 60.00, 0, None, None, "Bin 4"),
    ("A-009", "Mid-Century Teak Bowl", 8, "Estate sale", date(2026, 6, 14), 7.00, 2, 55.00, 1, None, None, "Shelf C"),
    ("A-010", "Patagonia Better Sweater M", 0, "Savers", date(2026, 6, 28), 12.00, 0, 64.00, 2, date(2026, 7, 9), 64.00, "Bin 2"),
    ("A-011", "Fossil Men's Watch", 3, "Garage sale", date(2026, 7, 5), 8.00, 0, 45.00, 1, None, None, "Drawer 1"),
    ("A-012", "KitchenAid Mixer Bowl", 4, "Thrift store", date(2026, 7, 12), 10.00, 1, 40.00, 1, None, None, "Shelf C"),
    ("A-013", "Funko Pop Exclusive", 6, "Flea market", date(2026, 7, 20), 14.00, 4, 22.00, 3, None, None, "Bin 5"),
    ("A-014", "Dr. Martens 1460 – Size 8", 1, "Goodwill", date(2026, 7, 26), 18.00, 3, 85.00, 1, None, None, "Shelf A"),
    ("A-015", "Nintendo Switch Games (4)", 5, "Facebook", date(2026, 8, 3), 60.00, 1, 110.00, 2, date(2026, 8, 10), 110.00, "Shelf D"),
    ("A-016", "Anthropologie Maxi Dress", 0, "Thrift store", date(2026, 8, 11), 7.00, 3, 36.00, 1, None, None, "Bin 2"),
    ("A-017", "Vintage Band Tee – Nirvana", 8, "Estate sale", date(2026, 8, 19), 4.00, 2, 55.00, 2, date(2026, 9, 1), 55.00, "Bin 3"),
    ("A-018", "Pandora Charm Bracelet", 3, "Garage sale", date(2026, 8, 30), 15.00, 0, 70.00, 1, None, None, "Drawer 1"),
    ("A-019", "Carhartt Detroit Jacket", 0, "Bins outlet", date(2026, 9, 6), 18.00, 0, 135.00, 1, None, None, "Rack 1"),
    ("A-020", "Le Creuset Dutch Oven", 4, "Estate sale", date(2026, 9, 13), 60.00, 1, 180.00, 0, None, None, "Shelf C"),
    ("A-021", "Kindle Paperwhite", 5, "Goodwill", date(2026, 9, 20), 20.00, None, None, 0, None, None, "Shelf D"),
    ("A-022", "Vintage Pyrex Casserole", 4, "Thrift store", date(2026, 9, 27), 12.00, 2, 45.00, 0, None, None, "Shelf C"),
]
for i, s in enumerate(samples):
    r = FIRST + i
    sku, item, cat, src, bought, cost, plat, lst, status, sold_d, sold_p, loc = s
    iv[f"B{r}"], iv[f"C{r}"], iv[f"D{r}"], iv[f"E{r}"] = sku, item, CATEGORIES[cat], src
    iv[f"F{r}"], iv[f"G{r}"], iv[f"I{r}"] = bought, cost, lst
    iv[f"H{r}"] = PLATFORMS[plat] if plat is not None else None
    iv[f"J{r}"], iv[f"K{r}"], iv[f"L{r}"], iv[f"M{r}"] = STATUSES[status], sold_d, sold_p, loc
SAMPLE_LAST = FIRST + len(samples) - 1
iv["B6"].comment = note(f"Rows 6–{SAMPLE_LAST} are EXAMPLES. Select B6:M{SAMPLE_LAST} and press Delete to start fresh.\n"
                        f"Las filas 6–{SAMPLE_LAST} son EJEMPLOS. Selecciona B6:M{SAMPLE_LAST} y presiona Suprimir.")
list_validation(iv, f"D{FIRST}:D{LAST}", source=f"=Settings!$F${CAT_FIRST}:$F${CAT_LAST}", title="Category / Categoría")
list_validation(iv, f"H{FIRST}:H{LAST}", PLATFORMS, title="Platform / Plataforma")
list_validation(iv, f"J{FIRST}:J{LAST}", STATUSES, title="Status / Estado")
date_validation(iv, f"F{FIRST}:F{LAST}")
date_validation(iv, f"K{FIRST}:K{LAST}")
number_validation(iv, [f"G{FIRST}:G{LAST}", f"I{FIRST}:I{LAST}", f"L{FIRST}:L{LAST}"])
contains_rule(iv, f"Q{FIRST}:Q{LAST}", "⚠", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(iv, f"Q{FIRST}:Q{LAST}", "WATCH", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(iv, f"Q{FIRST}:Q{LAST}", "FRESH", Font(name=FONT, color=GREEN), fill(GREEN_LIGHT))
contains_rule(iv, f"Q{FIRST}:Q{LAST}", "SOLD", Font(name=FONT, color=TEAL_DARK), fill(TEAL_LIGHT))
contains_rule(iv, f"Q{FIRST}:Q{LAST}", "CLOSED", Font(name=FONT, italic=True, color="94A3B8"))
iv.freeze_panes = "D6"
iv.auto_filter.ref = f"B5:Q{LAST}"
iv.print_title_rows = "5:5"

# =====================================================================
# DASHBOARD
# =====================================================================
db = wb.create_sheet("Dashboard", 1)
banner(db, "N", "Dashboard  /  Panel de Control",
       "What you own, what it's worth and what's not moving. Updates automatically.  •  "
       "Qué tienes, cuánto vale y qué no se vende. Se actualiza solo.")
for col, w in {"B": 34, "C": 13, "D": 15, "E": 15, "F": 13, "G": 3,
               "H": 13, "I": 13, "J": 13, "K": 13, "L": 13, "M": 13, "N": 3}.items():
    db.column_dimensions[col].width = w


def ir(col):
    return f"Inventory!${col}${FIRST}:${col}${LAST}"


kpi_cards(db, 6, [
    ("B", "B", "ITEMS IN STOCK", "Artículos en Inventario", f"=SUM({ir('S')})", "#,##0",
     f"=SUM({ir('AA')})", '0" not listed yet / sin publicar"'),
    ("C", "D", "INVENTORY VALUE (COST)", "Valor del Inventario (Costo)", f"=ROUND(SUMIFS({ir('G')},{ir('S')},1),2)", CUR0,
     "=IF(B8=0,0,C8/B8)", '$#,##0.00" avg cost / costo prom."'),
    ("E", "F", "POTENTIAL REVENUE", "Ingresos Potenciales", f"=ROUND(SUMIFS({ir('I')},{ir('S')},1),2)", CUR0,
     f"=SUM({ir('AC')})", '$#,##0" potential profit / ganancia potencial"'),
    ("H", "J", "STALE ITEMS", "Artículos Estancados", f"=SUM({ir('Z')})", "#,##0",
     f"=SUMIFS({ir('G')},{ir('Z')},1)", '$#,##0" of cost tied up / de costo inmovilizado"'),
])
kpi_cards(db, 11, [
    ("B", "B", "ITEMS SOLD", "Artículos Vendidos", f"=SUM({ir('T')})", "#,##0",
     f"=SUMIFS({ir('L')},{ir('T')},1)", '$#,##0" in sales / en ventas"'),
    ("C", "D", "AVG. DAYS TO SELL", "Días Promedio para Vender", f"=IF(SUM({ir('Y')})=0,0,SUM({ir('X')})/SUM({ir('Y')}))", "0.0",
     "purchase → sale / compra → venta", None),
    ("E", "F", "SELL-THROUGH RATE", "Tasa de Venta", "=IF(B8+B13=0,0,B13/(B8+B13))", PCT,
     "sold ÷ (sold + in stock) / vendidos ÷ total", None),
    ("H", "J", "GROSS PROFIT (SOLD)", "Ganancia Bruta (Vendidos)", f"=SUM({ir('AB')})", CUR0,
     f"=IF(SUMIFS({ir('G')},{ir('T')},1)=0,0,H13/SUMIFS({ir('G')},{ir('T')},1))", '0%" return on cost / retorno sobre costo"'),
])

# ---- By category ----
section_title(db, "B17:F17", "Inventory by Category  /  Inventario por Categoría")
db.row_dimensions[18].height = 30
for col, text in zip("BCDEF", ["Category\nCategoría", "In Stock\nEn Inventario", "Value at Cost\nValor al Costo",
                               "Potential Revenue\nIngreso Potencial", "Sold\nVendidos"]):
    header_cell(db[f"{col}18"], text)
C0 = 19
n = CAT_LAST - CAT_FIRST + 1
UNCAT, TOT = C0 + n, C0 + n + 1
rows = []
for i in range(n):
    r = C0 + i
    s = CAT_FIRST + i
    rows.append([
        f'=IF(Settings!$F${s}="","",Settings!$F${s})',
        f'=IF($B{r}="",0,COUNTIFS({ir("S")},1,{ir("D")},$B{r}))',
        f'=IF($B{r}="",0,ROUND(SUMIFS({ir("G")},{ir("S")},1,{ir("D")},$B{r}),2))',
        f'=IF($B{r}="",0,ROUND(SUMIFS({ir("I")},{ir("S")},1,{ir("D")},$B{r}),2))',
        f'=IF($B{r}="",0,COUNTIFS({ir("T")},1,{ir("D")},$B{r}))',
    ])
rows.append(["Uncategorized / Sin categoría",
             f"=C{TOT}-SUM(C{C0}:C{UNCAT - 1})", f"=ROUND(D{TOT}-SUM(D{C0}:D{UNCAT - 1}),2)",
             f"=ROUND(E{TOT}-SUM(E{C0}:E{UNCAT - 1}),2)", f"=F{TOT}-SUM(F{C0}:F{UNCAT - 1})"])
fmts = {"C": "#,##0", "D": CUR_SIMPLE, "E": CUR_SIMPLE, "F": "#,##0"}
table_rows(db, C0, rows, "BCDEF", fmts, bold_first=False)
total_row(db, TOT, "BCDEF", ["TOTAL", "=B8", "=C8", "=E8", "=B13"], fmts)

# ---- Age buckets ----
A_TITLE = TOT + 3
section_title(db, f"B{A_TITLE}:F{A_TITLE}", "Unsold Items by Age  /  Artículos sin Vender por Antigüedad")
A_HEAD = A_TITLE + 1
db.row_dimensions[A_HEAD].height = 30
for col, text in zip("BCDEF", ["Age\nAntigüedad", "Items\nArtículos", "Value at Cost\nValor al Costo",
                               "Potential Revenue\nIngreso Potencial", "% of Items\n% de Artículos"]):
    header_cell(db[f"{col}{A_HEAD}"], text)
A0 = A_HEAD + 1
buckets = ["0–30 days / días", "31–60 days / días", "61–90 days / días", "90+ days / días", "No purchase date / Sin fecha"]
rows = []
for b, label in enumerate(buckets, start=1):
    r = A0 + b - 1
    rows.append([label, f"=COUNTIF({ir('V')},{b})", f"=ROUND(SUMIFS({ir('G')},{ir('V')},{b}),2)",
                 f"=ROUND(SUMIFS({ir('I')},{ir('V')},{b}),2)", f"=IF($C${A0 + 5}=0,0,C{r}/$C${A0 + 5})"])
afmts = {"C": "#,##0", "D": CUR_SIMPLE, "E": CUR_SIMPLE, "F": "0.0%"}
table_rows(db, A0, rows, "BCDEF", afmts)
total_row(db, A0 + 5, "BCDEF", ["IN STOCK / EN INVENTARIO", f"=SUM(C{A0}:C{A0 + 4})", f"=SUM(D{A0}:D{A0 + 4})",
                                f"=SUM(E{A0}:E{A0 + 4})", f"=IF(C{A0 + 5}=0,0,1)"], afmts)

# ---- Oldest unsold ----
O_TITLE = A0 + 8
section_title(db, f"B{O_TITLE}:M{O_TITLE}", "10 Oldest Unsold Items – reprice, relist or bundle  /  "
                                            "10 Artículos sin Vender más Antiguos – baja el precio o vuelve a publicar")
O_HEAD = O_TITLE + 1
db.row_dimensions[O_HEAD].height = 30
o_cols = [("B", "Item\nArtículo"), ("C", "SKU\nCódigo"), ("D", "Days Held\nDías"), ("E", "Cost\nCosto"),
          ("F", "List Price\nPrecio"), ("H", "Platform\nPlataforma"), ("J", "Location\nUbicación"),
          ("K", "Alert\nAlerta")]
for col, text in o_cols:
    header_cell(db[f"{col}{O_HEAD}"], text)
db.merge_cells(f"H{O_HEAD}:I{O_HEAD}")
db.merge_cells(f"K{O_HEAD}:M{O_HEAD}")
header_cell(db[f"G{O_HEAD}"], "")
for k in range(1, 11):
    r = O_HEAD + k
    key = f"LARGE({ir('W')},{k})"
    idx = f"MATCH({key},{ir('W')},0)"
    none = f"{key}<-1E+14"

    def pick(col, blank='""'):
        return f'=IFERROR(IF({none},{blank},INDEX({ir(col)},{idx})),{blank})'

    db[f"B{r}"] = pick("C", '"—"')
    db[f"C{r}"] = pick("B")
    db[f"D{r}"] = pick("N")
    db[f"E{r}"] = pick("G")
    db[f"F{r}"] = pick("I")
    db.merge_cells(f"H{r}:I{r}")
    db[f"H{r}"] = pick("H")
    db[f"J{r}"] = pick("M")
    db.merge_cells(f"K{r}:M{r}")
    db[f"K{r}"] = pick("Q")
    for col in "BCDEFGHIJKLM":
        c = db[f"{col}{r}"]
        c.border = BORDER
        c.font = font(10)
        c.fill = fill(TEAL_LIGHT if k % 2 else WHITE)
    db[f"D{r}"].number_format = "0"
    db[f"D{r}"].alignment = Alignment(horizontal="center")
    db[f"D{r}"].font = font(10, True, NAVY)
    db[f"E{r}"].number_format = CUR
    db[f"F{r}"].number_format = CUR
contains_rule(db, f"K{O_HEAD + 1}:K{O_HEAD + 10}", "⚠", Font(name=FONT, bold=True, color=RED))
contains_rule(db, f"K{O_HEAD + 1}:K{O_HEAD + 10}", "WATCH", Font(name=FONT, bold=True, color=AMBER))

# ---- Charts ----
bar = bar_chart("Inventory Value by Category / Valor por Categoría", horizontal=True)
bar.add_data(Reference(db, min_col=4, min_row=18, max_row=UNCAT), titles_from_data=True)
bar.set_categories(Reference(db, min_col=2, min_row=C0, max_row=UNCAT))
color_series(bar.series[0], TEAL, TEAL_DARK)
bar.dataLabels.numFmt = "$#,##0"
bar.x_axis.scaling.orientation = "maxMin"
bar.y_axis.crosses = "max"
bar.gapWidth = 40
bar.x_axis.txPr = axis_text(750)
bar.height, bar.width = 8.6, 17
db.add_chart(bar, "H17")

age = bar_chart("Unsold Items by Age / Antigüedad del Inventario", money=False)
age.add_data(Reference(db, min_col=3, min_row=A_HEAD, max_row=A0 + 4), titles_from_data=True)
age.set_categories(Reference(db, min_col=2, min_row=A0, max_row=A0 + 4))
color_series(age.series[0], NAVY, NAVY)
age.dataLabels.numFmt = "0"
age.x_axis.txPr = axis_text(750)
age.height, age.width = 5.2, 17
db.add_chart(age, f"H{A_TITLE}")

# =====================================================================
# QUICK START GUIDE
# =====================================================================
steps = [
    (f"Clear the example data: on Inventory select B6:M{SAMPLE_LAST} and press Delete. Never delete whole rows or the grey columns.",
     f"Borra los datos de ejemplo: en Inventory selecciona B6:M{SAMPLE_LAST} y presiona Suprimir. Nunca borres filas completas ni las columnas grises."),
    ("Check Settings: choose when an unsold item becomes 'Watch' (default 60 days) and 'Stale' (default 90 days). Rename or add categories if you like.",
     "Revisa Settings (Configuración): elige cuándo un artículo sin vender pasa a 'Vigilar' (60 días por defecto) y 'Estancado' (90 días). Puedes renombrar o agregar categorías."),
    ("Add every item the day you buy it: SKU (your own code), name, category, source, purchase date and cost. Write where you stored it in Location / Bin so you can find it fast.",
     "Agrega cada artículo el día que lo compras: código (SKU), nombre, categoría, fuente, fecha y costo. Escribe dónde lo guardaste en Ubicación para encontrarlo rápido."),
    ("When you list it, choose the Platform, enter the List Price and set Status to 'Listed / Publicado'.",
     "Cuando lo publiques, elige la Plataforma, escribe el Precio Publicado y cambia el Estado a 'Listed / Publicado'."),
    ("When it sells, set Status to 'Sold / Vendido' and enter the Sold Date and Sold Price. Donated or damaged items: pick that status – they leave your stock totals.",
     "Cuando se venda, cambia el Estado a 'Sold / Vendido' y escribe la Fecha y el Precio de Venta. Artículos donados o dañados: elige ese estado – salen de tu inventario."),
    ("Grey columns fill in automatically: Days Held, Potential Profit (list − cost), Gross Profit (sold − cost) and the Age Alert: FRESH, WATCH or STALE.",
     "Las columnas grises se llenan solas: Días, Ganancia Potencial (precio − costo), Ganancia Bruta (venta − costo) y la Alerta: RECIENTE, VIGILAR o ESTANCADO."),
    ("Open the Dashboard: items in stock, inventory value, potential revenue, stale items, items sold, average days to sell, sell-through rate and gross profit – plus the 10 oldest unsold items to reprice or relist. It's locked to protect the formulas.",
     "Abre el Dashboard (Panel): artículos en inventario, valor, ingresos potenciales, estancados, vendidos, días promedio para vender, tasa de venta y ganancia bruta – más los 10 artículos más antiguos para bajar precio o volver a publicar. Está bloqueado para proteger las fórmulas."),
    ("Weekly habit: sort by Age Alert, reprice STALE items, relist or bundle them, and take a quick photo of your bins.",
     "Hábito semanal: ordena por Alerta, baja el precio de los ESTANCADOS, vuelve a publicarlos o haz lotes."),
    ("Gross profit here is before platform fees and shipping. For true net profit per sale, use ProfitTrack Profit Tracker.",
     "La ganancia bruta aquí es antes de comisiones y envío. Para la ganancia neta real por venta, usa ProfitTrack Profit Tracker."),
    ("Google Sheets: upload to Google Drive → Open with Google Sheets → File → Save as Google Sheets. Tip: save a blank copy first and back up monthly.",
     "Google Sheets: súbelo a Google Drive → Abrir con Google Sheets → Archivo → Guardar como Hojas de cálculo de Google. Consejo: guarda una copia en blanco y respalda cada mes."),
]
faq = ("FAQ  /  Preguntas Frecuentes", [
    ("Q: Why do Days Held change every day?  A: Unsold items count days up to today, so the file stays current on its own.",
     "P: ¿Por qué cambian los días cada día?  R: Los artículos sin vender cuentan días hasta hoy, así el archivo se actualiza solo."),
    ("Q: How many items?  A: 1,000 rows are ready. Need more? Copy the last row down and extend the Inventory ranges on the Dashboard.",
     "P: ¿Cuántos artículos?  R: Hay 1,000 filas listas. ¿Necesitas más? Copia la última fila hacia abajo y amplía los rangos en el Panel."),
    ("Q: Why can't I type on the Dashboard?  A: It's locked so formulas can't be erased by accident. To unlock (no password): Excel → Review → Unprotect Sheet.",
     "P: ¿Por qué no puedo escribir en el Panel?  R: Está bloqueado para no borrar fórmulas por accidente. Para desbloquear (sin contraseña): Excel → Revisar → Desproteger hoja."),
])
build_guide(wb, steps, [faq])

finish(wb, db, OUT, "ProfitTrack Inventory — Reseller Inventory Tracker")
