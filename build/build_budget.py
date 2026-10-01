"""Build ProfitTrack Budget — Bilingual Monthly Budget Planner (English / Español)."""
import random
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font

from common import (AMBER, AMBER_LIGHT, BORDER, CUR, CUR0, CUR_SIMPLE, DATE_FMT, FONT, GREEN, GREEN_LIGHT, MONTHS,
                    NAVY, NAVY_2, PCT, RED, RED_LIGHT, SLATE, TEAL, TEAL_DARK, TEAL_LIGHT, WHITE, axis_text, banner,
                    bar_chart, build_cover, build_guide, color_series, contains_rule, date_validation, fill, finish,
                    font, header_cell, hide_helpers, kpi_cards, list_validation, note, number_validation,
                    section_title, setting_cell, small_title, style_grid, table_rows, total_row)

OUT = "product/ProfitTrack_Bilingual_Budget_Planner.xlsx"
T_FIRST, T_LAST = 6, 1005       # 1,000 transactions
B_FIRST, B_LAST = 7, 41         # Budget categories (income 7-11, expenses 12-41)
INC_LAST = 11
G_FIRST, G_LAST = 6, 25         # savings goals
D_FIRST, D_LAST = 6, 25         # debts

INCOME, EXPENSE = "Income / Ingreso", "Expense / Gasto"
NEEDS, WANTS, SAVE = "Needs / Necesidades", "Wants / Gustos", "Savings & Debt / Ahorro y Deuda"
GROUPS = [NEEDS, WANTS, SAVE]
PAYMENTS = ["Debit / Débito", "Credit / Crédito", "Cash / Efectivo", "Transfer / Transferencia", "Other / Otro"]

INCOME_CATS = [("Salary / Salario", 3400), ("Partner salary / Salario de pareja", 2200),
               ("Side income / Ingreso extra", 300), ("Other income / Otros ingresos", 0), (None, None)]
EXPENSE_CATS = [
    ("Rent / Mortgage – Renta / Hipoteca", NEEDS, 1650),
    ("Utilities / Servicios (luz, agua, gas)", NEEDS, 220),
    ("Groceries / Supermercado", NEEDS, 650),
    ("Transportation / Transporte", NEEDS, 180),
    ("Gas / Gasolina", NEEDS, 200),
    ("Insurance / Seguros", NEEDS, 260),
    ("Phone & Internet / Teléfono e Internet", NEEDS, 140),
    ("Health / Salud", NEEDS, 120),
    ("Childcare & School / Niños y Escuela", NEEDS, 300),
    ("Family support / Apoyo a la familia (remesas)", NEEDS, 250),
    ("Dining out / Comer fuera", WANTS, 220),
    ("Entertainment / Entretenimiento", WANTS, 100),
    ("Shopping & Clothes / Compras y Ropa", WANTS, 150),
    ("Subscriptions / Suscripciones", WANTS, 45),
    ("Personal care / Cuidado personal", WANTS, 60),
    ("Gifts & Donations / Regalos y Donaciones", WANTS, 80),
    ("Travel / Viajes", WANTS, 100),
    ("Pets / Mascotas", WANTS, 50),
    ("Debt payments / Pagos de deudas", SAVE, 575),
    ("Emergency fund / Fondo de emergencia", SAVE, 200),
    ("Savings goals / Metas de ahorro", SAVE, 150),
    ("Retirement / Retiro", SAVE, 100),
    ("Other / Otro", WANTS, 50),
]

wb = Workbook()
build_cover(
    wb, "Bilingual Budget Planner", "Planificador de Presupuesto Bilingüe",
    [
        ("Dashboard", "📊  Dashboard / Panel", "Month at a glance, budget vs. actual, 50/30/20, year overview / Tu mes de un vistazo"),
        ("Transactions", "🧾  Transactions / Movimientos", "Log income & spending in seconds / Registra ingresos y gastos en segundos"),
        ("Budget", "🎯  Budget / Presupuesto", "Your monthly plan by category / Tu plan mensual por categoría"),
        ("Savings Goals", "🐷  Savings Goals / Metas de Ahorro", "Progress bars & how much to save each month / Cuánto ahorrar cada mes"),
        ("Debts", "💳  Debts / Deudas", "Payoff date, interest, snowball & avalanche order / Fecha de pago e intereses"),
        ("Quick Start Guide", "🚀  Quick Start Guide / Guía Rápida", "Step-by-step setup in 5 minutes / Configuración paso a paso"),
    ],
    "💻 PC / Mac  •  📱 iPhone / Android  —  Excel • Google Sheets   |   Computadora y celular — Excel • Google Sheets",
    kicker="FAMILY FINANCES  •  FINANZAS FAMILIARES",
    platforms="English  +  Español  —  every heading, every instruction / cada título, cada instrucción",
)

# =====================================================================
# BUDGET (categories + monthly plan)
# =====================================================================
bu = wb.create_sheet("Budget")
banner(bu, "F", "Monthly Budget  /  Presupuesto Mensual",
       "Plan how much you expect to earn and spend each month. Rename categories or add your own.  •  "
       "Planea cuánto esperas ganar y gastar cada mes. Renombra categorías o agrega las tuyas.")
for col, w in {"B": 46, "C": 18, "D": 30, "E": 16, "F": 30}.items():
    bu.column_dimensions[col].width = w
bu.row_dimensions[6].height = 36
for col, text in zip("BCDEF", ["Category\nCategoría", "Type\nTipo", "Group (50/30/20)\nGrupo", "Monthly Budget\nPresupuesto Mensual", "Notes\nNotas"]):
    header_cell(bu[f"{col}6"], text)
rows = [(n, INCOME, None, v) for n, v in INCOME_CATS] + [(n, EXPENSE, g, v) for n, g, v in EXPENSE_CATS]
rows += [(None, EXPENSE, None, None)] * (B_LAST - B_FIRST + 1 - len(rows))
for i, (name, typ, grp, val) in enumerate(rows):
    r = B_FIRST + i
    is_inc = typ == INCOME
    for col in "BCDEF":
        c = bu[f"{col}{r}"]
        c.border = BORDER
        c.fill = fill("FEF9C3" if col in "BDEF" else ("DCFCE7" if is_inc else "F1F5F9"))
        c.font = font(10, color="1E3A8A" if col in "BDEF" else NAVY, bold=(col == "C"))
    bu[f"B{r}"] = name
    bu[f"C{r}"] = typ
    bu[f"D{r}"] = grp if not is_inc else "—"
    bu[f"E{r}"] = val
    bu[f"E{r}"].number_format = CUR
list_validation(bu, f"D{INC_LAST + 1}:D{B_LAST}", GROUPS, title="Group / Grupo")
number_validation(bu, [f"E{B_FIRST}:E{B_LAST}"])
r = B_LAST + 1
for label, f, color in (
    ("TOTAL PLANNED INCOME / INGRESO PLANEADO", f"=SUM(E{B_FIRST}:E{INC_LAST})", NAVY),
    ("TOTAL PLANNED SPENDING / GASTO PLANEADO", f"=SUM(E{INC_LAST + 1}:E{B_LAST})", NAVY),
    ("LEFT TO PLAN / POR ASIGNAR (aim for $0 / ideal $0)", f"=E{B_LAST + 1}-E{B_LAST + 2}", TEAL_DARK),
):
    bu[f"B{r}"] = label
    bu.merge_cells(f"B{r}:D{r}")
    bu[f"E{r}"] = f
    for col in "BE":
        bu[f"{col}{r}"].font = font(10, True, WHITE)
        bu[f"{col}{r}"].fill = fill(color)
        bu[f"{col}{r}"].border = BORDER
    bu[f"E{r}"].number_format = CUR
    r += 1
bu[f"B{r + 1}"] = ("Zero-based budget: give every dollar a job until 'Left to plan' is $0.  /  "
                   "Presupuesto base cero: asigna cada dólar hasta que 'Por asignar' sea $0.")
bu[f"B{r + 1}"].font = font(8, False, SLATE, italic=True)
bu["B7"].comment = note("Yellow cells are yours: rename categories, change groups and amounts. Rows 11 and 35–41 are free for your own categories.\n"
                        "Las celdas amarillas son tuyas: renombra categorías, cambia grupos y montos. Las filas 11 y 35–41 están libres.")
bu.freeze_panes = "C7"
CATS = f"Budget!$B${B_FIRST}:$B${B_LAST}"
TYPES = f"Budget!$C${B_FIRST}:$C${B_LAST}"
GRPS = f"Budget!$D${B_FIRST}:$D${B_LAST}"

# =====================================================================
# TRANSACTIONS
# =====================================================================
tx = wb.create_sheet("Transactions", 1)
banner(tx, "I", "Transactions  /  Movimientos",
       "One row per income or expense. Pick a category – the type and group fill in automatically.  •  "
       "Una fila por ingreso o gasto. Elige la categoría – el tipo y el grupo se llenan solos.")
t_headers = [
    ("B", "Date\nFecha", 13, "in"),
    ("C", "Description\nDescripción", 30, "in"),
    ("D", "Category\nCategoría", 40, "in"),
    ("E", "Amount\nMonto", 13, "in"),
    ("F", "Payment Method\nMétodo de Pago", 18, "in"),
    ("G", "Notes\nNotas", 22, "in"),
    ("H", "Type\nTipo", 17, "calc"),
    ("I", "Check\nRevisar", 30, "calc"),
]
style_grid(tx, t_headers, 5, T_FIRST, T_LAST, money="E", dates="B", centered="HI")
for r in range(T_FIRST, T_LAST + 1):
    tx[f"K{r}"] = f'=IF(D{r}="","",IFERROR(MATCH(D{r},{CATS},0),0))'
    tx[f"H{r}"] = f'=IF(OR(E{r}="",D{r}=""),"",IF(K{r}=0,"{EXPENSE}",INDEX({TYPES},K{r})))'
    tx[f"I{r}"] = (f'=IF(AND(B{r}="",E{r}=""),"",IF(NOT(ISNUMBER(B{r})),"⚠ ADD DATE / AGREGA FECHA",'
                   f'IF(E{r}="","⚠ ADD AMOUNT / AGREGA MONTO",IF(D{r}="","⚠ ADD CATEGORY / AGREGA CATEGORÍA",'
                   f'IF(K{r}=0,"● NEW CATEGORY? / ¿CATEGORÍA NUEVA?","✔ OK")))))')
    tx[f"L{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(E{r}),D{r}<>""),YEAR(B{r})*100+MONTH(B{r}),0)'
    tx[f"M{r}"] = f'=IF(L{r}=0,0,YEAR(B{r}))'
    tx[f"N{r}"] = f'=IF(L{r}=0,"",IF(K{r}=0,"{WANTS}",INDEX({GRPS},K{r})))'
hide_helpers(tx, 5, [("K", "category row"), ("L", "year-month"), ("M", "year"), ("N", "group")], T_FIRST, T_LAST)
tx["E5"].comment = note("Always a positive number – the category decides if it is income or an expense.\n"
                        "Siempre un número positivo – la categoría decide si es ingreso o gasto.")

# deterministic sample data: Jan–Sep 2026
rnd = random.Random(42)
samples = []
for m in range(1, 10):
    samples += [
        (date(2026, m, 1), "Paycheck / Cheque", "Salary / Salario", 1700, 4),
        (date(2026, m, 15), "Paycheck / Cheque", "Salary / Salario", 1700, 4),
        (date(2026, m, 5), "Paycheck / Cheque", "Partner salary / Salario de pareja", 2200, 4),
        (date(2026, m, 1), "Rent / Renta", "Rent / Mortgage – Renta / Hipoteca", 1650, 3),
        (date(2026, m, 8), "Electric + water / Luz y agua", "Utilities / Servicios (luz, agua, gas)", round(rnd.uniform(170, 260), 2), 0),
        (date(2026, m, 6), "Walmart", "Groceries / Supermercado", round(rnd.uniform(150, 210), 2), 0),
        (date(2026, m, 13), "Costco", "Groceries / Supermercado", round(rnd.uniform(160, 240), 2), 1),
        (date(2026, m, 21), "H-E-B", "Groceries / Supermercado", round(rnd.uniform(120, 220), 2), 0),
        (date(2026, m, 10), "Shell", "Gas / Gasolina", round(rnd.uniform(80, 120), 2), 1),
        (date(2026, m, 24), "Shell", "Gas / Gasolina", round(rnd.uniform(70, 110), 2), 1),
        (date(2026, m, 12), "Car insurance / Seguro", "Insurance / Seguros", 260, 0),
        (date(2026, m, 18), "Verizon", "Phone & Internet / Teléfono e Internet", 140, 1),
        (date(2026, m, 20), "Remitly – mamá", "Family support / Apoyo a la familia (remesas)", 250, 3),
        (date(2026, m, 14), "Restaurant / Restaurante", "Dining out / Comer fuera", round(rnd.uniform(90, 280), 2), 1),
        (date(2026, m, 3), "Netflix + Spotify", "Subscriptions / Suscripciones", 28.98, 1),
        (date(2026, m, 25), "Credit card payment / Pago tarjeta", "Debt payments / Pagos de deudas", 575, 3),
        (date(2026, m, 26), "Savings transfer / Ahorro", "Emergency fund / Fondo de emergencia", 200, 3),
        (date(2026, m, 16), "Daycare / Guardería", "Childcare & School / Niños y Escuela", 300, 0),
    ]
    if m % 2 == 0:
        samples.append((date(2026, m, 19), "Target", "Shopping & Clothes / Compras y Ropa", round(rnd.uniform(60, 260), 2), 1))
    if m in (3, 6, 9):
        samples.append((date(2026, m, 22), "Side job / Trabajo extra", "Side income / Ingreso extra", 450, 4))
    if m == 7:
        samples.append((date(2026, 7, 4), "Trip to Galveston / Viaje", "Travel / Viajes", 640, 1))
samples.sort(key=lambda s: s[0])
for i, (d, desc, cat, amt, pay) in enumerate(samples):
    r = T_FIRST + i
    tx[f"B{r}"], tx[f"C{r}"], tx[f"D{r}"], tx[f"E{r}"], tx[f"F{r}"] = d, desc, cat, amt, PAYMENTS[pay]
TX_SAMPLE_LAST = T_FIRST + len(samples) - 1
tx["B6"].comment = note(f"Rows 6–{TX_SAMPLE_LAST} are EXAMPLES. Select B6:G{TX_SAMPLE_LAST} and press Delete to start fresh.\n"
                        f"Las filas 6–{TX_SAMPLE_LAST} son EJEMPLOS. Selecciona B6:G{TX_SAMPLE_LAST} y presiona Suprimir.")
list_validation(tx, f"D{T_FIRST}:D{T_LAST}", source=f"={CATS}", title="Category / Categoría")
list_validation(tx, f"F{T_FIRST}:F{T_LAST}", PAYMENTS, title="Payment / Pago")
date_validation(tx, f"B{T_FIRST}:B{T_LAST}")
number_validation(tx, [f"E{T_FIRST}:E{T_LAST}"])
contains_rule(tx, f"H{T_FIRST}:H{T_LAST}", "Income", Font(name=FONT, bold=True, color=GREEN), fill(GREEN_LIGHT))
contains_rule(tx, f"I{T_FIRST}:I{T_LAST}", "⚠", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(tx, f"I{T_FIRST}:I{T_LAST}", "NEW", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(tx, f"I{T_FIRST}:I{T_LAST}", "OK", Font(name=FONT, color=GREEN))
tx.freeze_panes = "C6"
tx.auto_filter.ref = f"B5:I{T_LAST}"
tx.print_title_rows = "5:5"

# =====================================================================
# SAVINGS GOALS
# =====================================================================
sg = wb.create_sheet("Savings Goals")
banner(sg, "J", "Savings Goals  /  Metas de Ahorro",
       "Set a target and a date – see your progress and how much to save each month.  •  "
       "Pon una meta y una fecha – mira tu avance y cuánto ahorrar cada mes.")
g_headers = [
    ("B", "Goal\nMeta", 28, "in"),
    ("C", "Target $\nMeta $", 13, "in"),
    ("D", "Saved So Far $\nAhorrado $", 14, "in"),
    ("E", "Target Date\nFecha Meta", 13, "in"),
    ("F", "Still Needed $\nFalta $", 13, "calc"),
    ("G", "Progress\nAvance", 10, "calc"),
    ("H", "Progress Bar\nBarra de Avance", 22, "calc"),
    ("I", "Months Left\nMeses", 10, "calc"),
    ("J", "Save per Month\nAhorro Mensual", 15, "calc"),
]
style_grid(sg, g_headers, 5, G_FIRST, G_LAST, money="CDFJ", pct="G", dates="E", centered="HI")
for r in range(G_FIRST, G_LAST + 1):
    sg[f"F{r}"] = f'=IF(OR(B{r}="",C{r}=""),"",MAX(0,C{r}-IF(D{r}="",0,D{r})))'
    sg[f"G{r}"] = f'=IF(F{r}="","",IF(C{r}=0,0,MIN(1,IF(D{r}="",0,D{r})/C{r})))'
    sg[f"H{r}"] = f'=IF(G{r}="","",REPT("█",ROUND(G{r}*10,0))&REPT("░",10-ROUND(G{r}*10,0)))'
    sg[f"H{r}"].font = font(11, False, TEAL_DARK)
    sg[f"I{r}"] = (f'=IF(OR(F{r}="",NOT(ISNUMBER(E{r}))),"",'
                   f'MAX(1,(YEAR(E{r})-YEAR(TODAY()))*12+MONTH(E{r})-MONTH(TODAY())))')
    sg[f"J{r}"] = f'=IF(OR(F{r}="",I{r}=""),"",F{r}/I{r})'
    sg[f"L{r}"] = f'=IF(F{r}="",0,1)'
hide_helpers(sg, 5, [("L", "goal")], G_FIRST, G_LAST)
goals = [("Emergency fund / Fondo de emergencia", 3000, 1250, date(2027, 6, 30)),
         ("Trip to Mexico / Viaje a México", 1800, 600, date(2027, 3, 1)),
         ("Car down payment / Enganche del carro", 5000, 900, date(2028, 1, 1)),
         ("Christmas gifts / Regalos de Navidad", 600, 250, date(2026, 12, 15))]
for i, (g, t, s_, d) in enumerate(goals):
    r = G_FIRST + i
    sg[f"B{r}"], sg[f"C{r}"], sg[f"D{r}"], sg[f"E{r}"] = g, t, s_, d
sg["I5"].comment = note("Counted from this month to the target month (minimum 1).\nContado desde este mes hasta el mes meta (mínimo 1).")
date_validation(sg, f"E{G_FIRST}:E{G_LAST}")
number_validation(sg, [f"C{G_FIRST}:D{G_LAST}"])
sg.freeze_panes = "C6"

# =====================================================================
# DEBTS
# =====================================================================
de = wb.create_sheet("Debts")
banner(de, "L", "Debt Payoff  /  Pago de Deudas",
       "See when each debt will be paid off and how much interest it costs. Extra payments make a big difference.  •  "
       "Mira cuándo terminas de pagar cada deuda y cuánto interés pagas. Los pagos extra hacen gran diferencia.")
d_headers = [
    ("B", "Debt\nDeuda", 26, "in"),
    ("C", "Balance $\nSaldo $", 13, "in"),
    ("D", "APR %\nInterés Anual", 12, "in"),
    ("E", "Minimum Payment\nPago Mínimo", 14, "in"),
    ("F", "Extra Payment\nPago Extra", 13, "in"),
    ("G", "Monthly Payment\nPago Mensual", 14, "calc"),
    ("H", "Months to Pay Off\nMeses para Pagar", 14, "calc"),
    ("I", "Payoff Date\nFecha de Pago", 13, "calc"),
    ("J", "Total Interest\nInterés Total", 13, "calc"),
    ("K", "Snowball Order\nOrden Bola de Nieve", 14, "calc"),
    ("L", "Avalanche Order\nOrden Avalancha", 14, "calc"),
]
style_grid(de, d_headers, 5, D_FIRST, D_LAST, money="CEFGJ", pct="D", dates="I", centered="HKL")
for r in range(D_FIRST, D_LAST + 1):
    de[f"D{r}"].number_format = "0.00%"
    pay = f"G{r}"
    nper = f'NPER(IF(D{r}="",0,D{r})/12,-{pay},C{r})'
    de[f"G{r}"] = f'=IF(OR(B{r}="",C{r}=""),"",IF(E{r}="",0,E{r})+IF(F{r}="",0,F{r}))'
    de[f"H{r}"] = (f'=IF(G{r}="","",IF(C{r}<=0,0,IF(G{r}<=C{r}*IF(D{r}="",0,D{r})/12,"⚠ Payment too low / Pago muy bajo",'
                   f'ROUNDUP({nper},0))))')
    de[f"I{r}"] = f'=IF(OR(H{r}="",NOT(ISNUMBER(H{r}))),"",EDATE(TODAY(),H{r}))'
    de[f"J{r}"] = f'=IF(OR(H{r}="",NOT(ISNUMBER(H{r}))),"",IF(H{r}=0,0,MAX(0,{pay}*{nper}-C{r})))'
    de[f"N{r}"] = f'=IF(OR(B{r}="",C{r}="",C{r}<=0),0,1)'
    de[f"O{r}"] = f'=IF(N{r}=1,C{r}+ROW()/1000000,"")'
    de[f"P{r}"] = f'=IF(N{r}=1,IF(D{r}="",0,D{r})-ROW()/1000000,"")'
    de[f"K{r}"] = f'=IF(N{r}=0,"",COUNTIF($O${D_FIRST}:$O${D_LAST},"<"&O{r})+1)'
    de[f"L{r}"] = f'=IF(N{r}=0,"",COUNTIF($P${D_FIRST}:$P${D_LAST},">"&P{r})+1)'
hide_helpers(de, 5, [("N", "active"), ("O", "snowball key"), ("P", "avalanche key")], D_FIRST, D_LAST)
debts = [("Credit card / Tarjeta de crédito", 2400, 0.2499, 75, 50),
         ("Car loan / Préstamo del carro", 9800, 0.075, 310, 0),
         ("Medical bill / Cuenta médica", 650, 0, 50, 0),
         ("Student loan / Préstamo estudiantil", 12000, 0.055, 140, 0)]
for i, (n, b, apr, mn, ex) in enumerate(debts):
    r = D_FIRST + i
    de[f"B{r}"], de[f"C{r}"], de[f"D{r}"], de[f"E{r}"], de[f"F{r}"] = n, b, apr, mn, ex
de["K5"].comment = note("Snowball: pay the SMALLEST balance first (quick wins).\nBola de nieve: paga primero el saldo MÁS PEQUEÑO.")
de["L5"].comment = note("Avalanche: pay the HIGHEST interest first (saves the most money).\nAvalancha: paga primero el interés MÁS ALTO (ahorras más).")
de["J5"].comment = note("Estimate at your current monthly payment.\nEstimado con tu pago mensual actual.")
number_validation(de, [f"C{D_FIRST}:C{D_LAST}", f"E{D_FIRST}:F{D_LAST}"])
number_validation(de, [f"D{D_FIRST}:D{D_LAST}"], high="1", title="APR", error="Enter the yearly rate, e.g. 24.99%.\nEscribe la tasa anual, ej. 24.99%.")
contains_rule(de, f"H{D_FIRST}:H{D_LAST}", "⚠", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(de, f"K{D_FIRST}:L{D_LAST}", "1", Font(name=FONT, bold=True, color=TEAL_DARK))
r = D_LAST + 2
for label, f, fmt in (("TOTAL DEBT / DEUDA TOTAL", f"=SUM(C{D_FIRST}:C{D_LAST})", CUR),
                      ("TOTAL MONTHLY PAYMENTS / PAGOS MENSUALES", f"=SUM(G{D_FIRST}:G{D_LAST})", CUR),
                      ("TOTAL INTEREST (estimate) / INTERÉS TOTAL", f"=SUM(J{D_FIRST}:J{D_LAST})", CUR)):
    de[f"B{r}"] = label
    de.merge_cells(f"B{r}:F{r}")
    de[f"G{r}"] = f
    for col in "BG":
        de[f"{col}{r}"].font = font(10, True, WHITE)
        de[f"{col}{r}"].fill = fill(NAVY)
        de[f"{col}{r}"].border = BORDER
    de[f"G{r}"].number_format = fmt
    r += 1
DEBT_TOTAL = f"Debts!$G${D_LAST + 2}"
de.freeze_panes = "C6"

# =====================================================================
# DASHBOARD
# =====================================================================
db = wb.create_sheet("Dashboard", 1)
banner(db, "O", "Dashboard  /  Panel de Control",
       "Pick a month to see where your money went. Updates automatically.  •  Elige un mes para ver a dónde fue tu dinero. Se actualiza solo.")
for col, w in {"B": 40, "C": 13, "D": 13, "E": 13, "F": 11, "G": 26, "H": 3,
               "I": 24, "J": 13, "K": 13, "L": 11, "M": 13, "N": 13, "O": 3}.items():
    db.column_dimensions[col].width = w


def tr(col):
    return f"Transactions!${col}${T_FIRST}:${col}${T_LAST}"


A0 = 58                                   # annual table first data row (months)
MONTH_LABELS = f"$B${A0}:$B${A0 + 11}"
db["B6"] = "Month / Mes:"
db["B6"].font = font(11, True, NAVY)
db["B6"].alignment = Alignment(horizontal="right", vertical="center")
setting_cell(db["C6"], f'=INDEX({MONTH_LABELS},IF(COUNT({tr("B")})=0,MONTH(TODAY()),MONTH(MAX({tr("B")}))))')
db.merge_cells("C6:D6")
db["E6"] = "Year / Año:"
db["E6"].font = font(11, True, NAVY)
db["E6"].alignment = Alignment(horizontal="right", vertical="center")
setting_cell(db["F6"], f'=IF(COUNT({tr("B")})=0,YEAR(TODAY()),YEAR(MAX({tr("B")})))', "0")
db["G6"] = "← pick month & year / elige mes y año"
db["G6"].font = font(8, False, SLATE, italic=True)
db.row_dimensions[6].height = 24
list_validation(db, "C6", MONTHS, title="Month / Mes")
db["C6"].comment = note("Shows the month of your latest transaction. Pick any month from the list.\n"
                        "Muestra el mes de tu último movimiento. Elige cualquier mes de la lista.")
db["Q6"] = f"=IFERROR(MATCH($C$6,{MONTH_LABELS},0),1)"
db["Q7"] = "=$F$6*100+$Q$6"
db.column_dimensions["Q"].hidden = True
KEY, YEAR = "$Q$7", "$F$6"

kpi_cards(db, 8, [
    ("B", "B", "INCOME THIS MONTH", "Ingresos del Mes", f'=SUMIFS({tr("E")},{tr("L")},{KEY},{tr("H")},"{INCOME}")', CUR0,
     f"=SUM(Budget!$E${B_FIRST}:$E${INC_LAST})", '$#,##0" planned / planeado"'),
    ("C", "D", "SPENDING THIS MONTH", "Gastos del Mes", f'=SUMIFS({tr("E")},{tr("L")},{KEY},{tr("H")},"{EXPENSE}")', CUR0,
     "=SUM(C21:C50)", '$#,##0" budgeted / presupuestado"'),
    ("E", "F", "LEFT OVER", "Sobrante", "=B10-C10", CUR0,
     "=IF(B10=0,0,E10/B10)", '0%" of income / del ingreso";[Red]-0%" of income / del ingreso"'),
    ("G", "G", "CATEGORIES OVER BUDGET", "Categorías Excedidas", "=SUM($S$21:$S$50)", "0",
     "✔ under = green · ⚠ over = red", None),
])
kpi_cards(db, 13, [
    ("B", "B", "INCOME THIS YEAR", "Ingresos del Año", f'=SUMIFS({tr("E")},{tr("M")},{YEAR},{tr("H")},"{INCOME}")', CUR0, None, None),
    ("C", "D", "SPENDING THIS YEAR", "Gastos del Año", f'=SUMIFS({tr("E")},{tr("M")},{YEAR},{tr("H")},"{EXPENSE}")', CUR0, None, None),
    ("E", "F", "SAVED TOWARD GOALS", "Ahorrado para Metas", "=SUM('Savings Goals'!$D$6:$D$25)", CUR0,
     "=SUM('Savings Goals'!$C$6:$C$25)", '"of "$#,##0" goals / de metas"'),
    ("G", "G", "TOTAL DEBT", "Deuda Total", f"={DEBT_TOTAL}", CUR0, f"=Debts!$G${D_LAST + 3}", '$#,##0"/mo in payments · pagos al mes"'),
])

# ---- Budget vs Actual (selected month) ----
section_title(db, "B18:G18", "Budget vs. Actual – Selected Month  /  Presupuesto vs. Real – Mes Elegido")
db.row_dimensions[19].height = 4
db.row_dimensions[20].height = 30
for col, text in zip("BCDEFG", ["Category\nCategoría", "Budget\nPresupuesto", "Actual\nReal", "Difference\nDiferencia",
                                "% Used\n% Usado", "Status\nEstado"]):
    header_cell(db[f"{col}20"], text)
BV0 = 21
rows = []
for i in range(B_LAST - INC_LAST):
    r = BV0 + i
    s_ = INC_LAST + 1 + i
    rows.append([
        f'=IF(Budget!$B${s_}="","",Budget!$B${s_})',
        f'=IF($B{r}="",0,IF(Budget!$E${s_}="",0,Budget!$E${s_}))',
        f'=IF($B{r}="",0,SUMIFS({tr("E")},{tr("L")},{KEY},{tr("D")},$B{r}))',
        f"=C{r}-D{r}",
        f'=IF($B{r}="","",IF(C{r}=0,IF(D{r}>0,1,0),D{r}/C{r}))',
        f'=IF($B{r}="","",IF(D{r}>C{r},"⚠ OVER / EXCEDIDO",IF(D{r}=0,"– not used / sin uso","✔ UNDER / BAJO")))',
    ])
    db[f"S{r}"] = f'=IF($B{r}="",0,IF(D{r}>C{r},1,0))'
db.column_dimensions["S"].hidden = True
bfmts = {"C": CUR, "D": CUR, "E": CUR, "F": "0%"}
table_rows(db, BV0, rows, "BCDEFG", bfmts, bold_first=False)
BV_LAST = BV0 + len(rows) - 1          # 50
UNC = BV_LAST + 1
unc_actual = f'=C10-SUM(D{BV0}:D{BV_LAST})'
table_rows(db, UNC, [["Other categories / Otras categorías", 0, unc_actual, f"=C{UNC}-D{UNC}", "", ""]], "BCDEFG", bfmts, bold_first=False)
total_row(db, UNC + 1, "BCDEFG", ["TOTAL SPENDING / GASTO TOTAL", f"=SUM(C{BV0}:C{UNC})", f"=SUM(D{BV0}:D{UNC})",
                                   f"=C{UNC + 1}-D{UNC + 1}", f"=IF(C{UNC + 1}=0,0,D{UNC + 1}/C{UNC + 1})", ""], bfmts)
contains_rule(db, f"G{BV0}:G{BV_LAST}", "OVER", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(db, f"G{BV0}:G{BV_LAST}", "UNDER", Font(name=FONT, color=GREEN), fill(GREEN_LIGHT))
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
db.conditional_formatting.add(f"E{BV0}:E{UNC + 1}", CellIsRule(operator="lessThan", formula=["0"],
                                                               font=Font(name=FONT, bold=True, color=RED)))

# ---- 50/30/20 (selected month) ----
section_title(db, "I18:N18", "50 / 30 / 20 Check – Selected Month  /  Regla 50/30/20")
db.row_dimensions[20].height = 30
for col, text in zip("IJKL", ["Group\nGrupo", "Spent\nGastado", "% of Income\n% del Ingreso", "Ideal\nIdeal"]):
    header_cell(db[f"{col}20"], text)
db.merge_cells("M20:N20")
header_cell(db["M20"], "Result\nResultado")
for i, (g, ideal) in enumerate(((NEEDS, 0.5), (WANTS, 0.3), (SAVE, 0.2))):
    r = 21 + i
    db[f"I{r}"] = g
    db[f"J{r}"] = f'=ROUND(SUMIFS({tr("E")},{tr("L")},{KEY},{tr("N")},I{r},{tr("H")},"{EXPENSE}"),2)'
    db[f"K{r}"] = f"=IF($B$10=0,0,J{r}/$B$10)"
    db[f"L{r}"] = ideal
    db.merge_cells(f"M{r}:N{r}")
    if g == SAVE:
        db[f"M{r}"] = f'=IF(K{r}>=L{r},"✔ GREAT / EXCELENTE","● SAVE MORE / AHORRA MÁS")'
    else:
        db[f"M{r}"] = f'=IF(K{r}<=L{r},"✔ ON TRACK / BIEN","● TOO HIGH / MUY ALTO")'
    for col in "IJKLMN":
        c = db[f"{col}{r}"]
        c.border = BORDER
        c.font = font(10, bold=(col == "I"))
        c.fill = fill(TEAL_LIGHT if i % 2 == 0 else WHITE)
    db[f"J{r}"].number_format = CUR_SIMPLE
    db[f"K{r}"].number_format = "0%"
    db[f"L{r}"].number_format = "0%"
contains_rule(db, "M21:M23", "●", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(db, "M21:M23", "✔", Font(name=FONT, bold=True, color=GREEN), fill(GREEN_LIGHT))
db["I24"] = "Groups come from the Budget sheet. / Los grupos vienen de la hoja Presupuesto."
db["I24"].font = font(8, False, SLATE, italic=True)

grp = BarChart()
grp.type = "col"
grp.title = small_title("Spending vs. Ideal / Gasto vs. Ideal (% of income)")
grp.style = 10
grp.add_data(Reference(db, min_col=11, min_row=20, max_row=23), titles_from_data=True)
grp.add_data(Reference(db, min_col=12, min_row=20, max_row=23), titles_from_data=True)
grp.set_categories(Reference(db, min_col=9, min_row=21, max_row=23))
grp.series[0].tx = SeriesLabel(v="Actual / Real")
grp.series[1].tx = SeriesLabel(v="Ideal")
color_series(grp.series[0], TEAL, TEAL_DARK)
color_series(grp.series[1], "CBD5E1", "94A3B8")
grp.y_axis.numFmt = "0%"
grp.y_axis.majorGridlines = None
grp.x_axis.delete = False
grp.y_axis.delete = False
grp.x_axis.txPr = axis_text(750)
grp.legend.position = "t"
grp.height, grp.width = 7.5, 17
db.add_chart(grp, "I26")

# ---- Annual overview ----
A_TITLE = A0 - 3
section_title(db, f"B{A_TITLE}:G{A_TITLE}", "Year Overview  /  Resumen del Año")
db.row_dimensions[A0 - 1].height = 30
for col, text in zip("BCDEF", ["Month\nMes", "Income\nIngresos", "Spending\nGastos", "Left Over\nSobrante", "Saved %\n% Ahorro"]):
    header_cell(db[f"{col}{A0 - 1}"], text)
rows = []
for m, label in enumerate(MONTHS, start=1):
    r = A0 + m - 1
    key = f"{YEAR}*100+{m}"
    rows.append([label, f'=ROUND(SUMIFS({tr("E")},{tr("L")},{key},{tr("H")},"{INCOME}"),2)',
                 f'=ROUND(SUMIFS({tr("E")},{tr("L")},{key},{tr("H")},"{EXPENSE}"),2)', f"=C{r}-D{r}",
                 f"=IF(C{r}=0,0,E{r}/C{r})"])
afmts = {"C": CUR_SIMPLE, "D": CUR_SIMPLE, "E": CUR, "F": PCT}
table_rows(db, A0, rows, "BCDEF", afmts)
total_row(db, A0 + 12, "BCDEF", ["YEAR / AÑO", "=B15", "=C15", f"=C{A0 + 12}-D{A0 + 12}",
                                 f"=IF(C{A0 + 12}=0,0,E{A0 + 12}/C{A0 + 12})"], afmts)
yr = BarChart()
yr.type = "col"
yr.title = small_title("Income vs. Spending / Ingresos vs. Gastos")
yr.style = 10
yr.add_data(Reference(db, min_col=3, min_row=A0 - 1, max_row=A0 + 11), titles_from_data=True)
yr.add_data(Reference(db, min_col=4, min_row=A0 - 1, max_row=A0 + 11), titles_from_data=True)
yr.set_categories(Reference(db, min_col=2, min_row=A0, max_row=A0 + 11))
yr.series[0].tx = SeriesLabel(v="Income / Ingresos")
yr.series[1].tx = SeriesLabel(v="Spending / Gastos")
color_series(yr.series[0], TEAL, TEAL_DARK)
color_series(yr.series[1], NAVY_2, NAVY)
yr.y_axis.numFmt = "$#,##0"
yr.y_axis.majorGridlines = None
yr.x_axis.delete = False
yr.y_axis.delete = False
yr.x_axis.txPr = axis_text(700)
yr.legend.position = "t"
yr.height, yr.width = 8.2, 17
db.add_chart(yr, f"I{A_TITLE}")

# =====================================================================
# QUICK START GUIDE
# =====================================================================
steps = [
    (f"Clear the example data: on Transactions select B6:G{TX_SAMPLE_LAST} and press Delete. Replace the example goals and debts with yours.",
     f"Borra los datos de ejemplo: en Transactions selecciona B6:G{TX_SAMPLE_LAST} y presiona Suprimir. Reemplaza las metas y deudas de ejemplo con las tuyas."),
    ("Budget: enter your monthly income, then give every dollar a job across the categories until 'Left to plan' is $0. Rename categories or add your own in the empty yellow rows.",
     "Budget (Presupuesto): escribe tu ingreso mensual y asigna cada dólar a las categorías hasta que 'Por asignar' sea $0. Renombra categorías o agrega las tuyas en las filas amarillas vacías."),
    ("Each category belongs to a group: Needs, Wants or Savings & Debt. This powers the 50/30/20 check on the Dashboard.",
     "Cada categoría pertenece a un grupo: Necesidades, Gustos o Ahorro y Deuda. Esto alimenta la regla 50/30/20 del Panel."),
    ("Transactions: log every paycheck and every expense – date, description, category (dropdown) and amount (always positive). The type fills in automatically.",
     "Transactions (Movimientos): registra cada cheque y cada gasto – fecha, descripción, categoría (lista) y monto (siempre positivo). El tipo se llena solo."),
    ("The Check column flags a missing date, amount or category, and categories that aren't in your Budget yet.",
     "La columna Revisar marca si falta fecha, monto o categoría, y categorías que todavía no están en tu Presupuesto."),
    ("Dashboard: pick a month and year (yellow cells). See income, spending, money left over, budget vs. actual for every category, the 50/30/20 check and your whole year.",
     "Dashboard (Panel): elige mes y año (celdas amarillas). Mira ingresos, gastos, sobrante, presupuesto vs. real por categoría, la regla 50/30/20 y todo tu año."),
    ("Savings Goals: add a goal, target amount, what you've saved and a target date – get a progress bar and how much to save each month.",
     "Savings Goals (Metas): agrega una meta, el monto, lo que llevas ahorrado y una fecha – verás una barra de avance y cuánto ahorrar cada mes."),
    ("Debts: enter balance, APR and payments – see months to pay off, payoff date and total interest, plus Snowball and Avalanche payoff order.",
     "Debts (Deudas): escribe saldo, interés anual y pagos – verás meses para pagar, fecha de pago e interés total, más el orden Bola de Nieve y Avalancha."),
    ("Weekly habit: 10 minutes every Sunday to log receipts and check the Dashboard. Small habits, big results!",
     "Hábito semanal: 10 minutos cada domingo para registrar recibos y revisar el Panel. ¡Pequeños hábitos, grandes resultados!"),
    ("Google Sheets: upload to Google Drive → Open with Google Sheets → File → Save as Google Sheets. Share it with your partner to budget together.",
     "Google Sheets: súbelo a Google Drive → Abrir con Google Sheets → Archivo → Guardar como Hojas de cálculo de Google. Compártelo con tu pareja para planear juntos."),
]
faq = ("FAQ  /  Preguntas Frecuentes", [
    ("Q: What is 50/30/20?  A: A simple guideline: about 50% of income for needs, 30% for wants and 20% for savings & debt payoff. Adjust it to your life.",
     "P: ¿Qué es 50/30/20?  R: Una guía sencilla: cerca del 50% del ingreso para necesidades, 30% para gustos y 20% para ahorro y deudas. Ajústala a tu vida."),
    ("Q: Snowball or avalanche?  A: Snowball pays the smallest balance first for quick motivation; avalanche pays the highest interest first to save the most money.",
     "P: ¿Bola de nieve o avalancha?  R: Bola de nieve paga primero el saldo más pequeño (motivación rápida); avalancha paga primero el interés más alto (ahorras más)."),
    ("Q: Why can't I type on the Dashboard?  A: It's locked so formulas can't be erased. Only the yellow month and year cells change. To unlock (no password): Excel → Review → Unprotect Sheet.",
     "P: ¿Por qué no puedo escribir en el Panel?  R: Está bloqueado para no borrar fórmulas. Solo cambian las celdas amarillas de mes y año. Para desbloquear (sin contraseña): Excel → Revisar → Desproteger hoja."),
    ("Q: Another currency?  A: Amounts show a $ sign but the math works for any currency – pesos, quetzales, euros…",
     "P: ¿Otra moneda?  R: Los montos muestran $ pero el cálculo sirve para cualquier moneda – pesos, quetzales, euros…"),
])
build_guide(wb, steps, [faq])

finish(wb, db, OUT, "ProfitTrack Budget — Bilingual Monthly Budget Planner")
