"""Render real screenshots of workbook ranges (LibreOffice -> PDF -> PNG, auto-cropped)."""
import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.worksheet.page import PageMargins
from PIL import Image, ImageChops

OUT_DIR = "build/shots"

SHOTS = {
    "product/ProfitTrack_Reseller_Profit_Tracker.xlsx": [
        ("pt_cover", "Cover", "A1:L38"),
        ("pt_dash", "Dashboard", "A1:T46"),
        ("pt_dash_top", "Dashboard", "A5:T36"),
        ("pt_kpis", "Dashboard", "B6:L9"),
        ("pt_top5", "Dashboard", "B39:L46"),
        ("pt_tracker", "Tracker", "A1:O25"),
        ("pt_tracker_calc", "Tracker", "G5:O25"),
        ("pt_phone", "Tracker", "B5:C30"),
        ("pt_sourcing", "Sourcing List", "A1:Q15"),
        ("pt_sourcing_calc", "Sourcing List", "B10:P15"),
        ("pt_sourcing_in", "Sourcing List", "B10:G15"),
        ("pt_sourcing_out", "Sourcing List", "L10:P15"),
        ("pt_guide", "Quick Start Guide", "A1:D11"),
        ("pt_guide_strip", "Quick Start Guide", "A1:D7"),
    ],
    "product/ProfitTrack_Expense_Mileage_Tracker.xlsx": [
        ("ex_cover", "Cover", "A1:L40"),
        ("ex_dash", "Dashboard", "A1:R48"),
        ("ex_kpis", "Dashboard", "B6:J11"),
        ("ex_expenses", "Expenses", "A1:L23"),
        ("ex_mileage", "Mileage", "A1:M21"),
        ("ex_mileage_calc", "Mileage", "B5:M21"),
        ("ex_phone", "Expenses", "B5:D23"),
        ("ex_settings", "Settings", "A1:F21"),
        ("ex_guide", "Quick Start Guide", "A1:D11"),
        ("ex_guide_strip", "Quick Start Guide", "A1:D7"),
    ],
    "product/ProfitTrack_Inventory_Tracker.xlsx": [
        ("iv_cover", "Cover", "A1:L37"),
        ("iv_dash", "Dashboard", "A1:R56"),
        ("iv_dash_top", "Dashboard", "A5:R43"),
        ("iv_kpis", "Dashboard", "B6:J14"),
        ("iv_oldest", "Dashboard", "B45:M56"),
        ("iv_inventory", "Inventory", "A1:Q27"),
        ("iv_alerts", "Inventory", "I5:Q27"),
        ("iv_phone", "Inventory", "C5:D27"),
        ("iv_guide", "Quick Start Guide", "A1:D11"),
        ("iv_guide_strip", "Quick Start Guide", "A1:D7"),
    ],
}


def trim(img, pad=0):
    bg = Image.new(img.mode, img.size, (255, 255, 255))
    box = ImageChops.difference(img, bg).getbbox()
    if box:
        img = img.crop((max(0, box[0] - pad), max(0, box[1] - pad),
                        min(img.width, box[2] + pad), min(img.height, box[3] + pad)))
    return img


def render(xlsx, name, sheet, rng, tmp, dpi):
    wb = load_workbook(xlsx)
    target = wb[sheet]
    wb.active = wb.sheetnames.index(sheet)
    for ws in wb.worksheets:
        if ws.title != sheet:
            ws.sheet_state = "hidden"
    target.print_area = rng
    target.print_title_rows = None
    target.page_margins = PageMargins(left=0, right=0, top=0, bottom=0, header=0, footer=0)
    target.page_setup.fitToWidth = 1
    target.page_setup.fitToHeight = 1
    target.sheet_properties.pageSetUpPr.fitToPage = True
    target.oddHeader.center.text = ""
    target.oddFooter.center.text = ""
    path = os.path.join(tmp, f"{name}.xlsx")
    wb.save(path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                   check=True, capture_output=True, timeout=240)
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", "-singlefile", "-f", "1", "-l", "1",
                    os.path.join(tmp, f"{name}.pdf"), os.path.join(tmp, name)], check=True)
    img = Image.open(os.path.join(tmp, f"{name}.png")).convert("RGB")
    trim(img).save(os.path.join(OUT_DIR, f"{name}.png"))
    print("shot", name, Image.open(os.path.join(OUT_DIR, f"{name}.png")).size)


if __name__ == "__main__":
    only = set(sys.argv[1:])
    os.makedirs(OUT_DIR, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        for xlsx, shots in SHOTS.items():
            for name, sheet, rng in shots:
                if not only or name in only:
                    render(xlsx, name, sheet, rng, tmp, dpi=300)
