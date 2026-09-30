"""Build ProfitTrack Expenses — Reseller Expense & Mileage Tracker (bilingual EN/ES)."""
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font

from common import (AMBER, AMBER_LIGHT, BORDER, CALC_FILL, CUR, CUR0, CUR_SIMPLE, DATE_FMT, FONT, GREEN,
                    GREEN_LIGHT, MONTHS, NAVY, NAVY_2, NO, RED, RED_LIGHT, SLATE, TEAL, TEAL_DARK, YES,
                    banner, bar_chart, build_cover, build_guide, color_series, contains_rule, date_validation,
                    fill, finish, font, header_cell, hide_helpers, kpi_cards, list_validation, note,
                    number_validation, section_title, setting_cell, small_title, axis_text, style_grid, table_rows, total_row)

OUT = "product/ProfitTrack_Expense_Mileage_Tracker.xlsx"
FIRST, LAST = 6, 505          # 500 expense rows / 500 trip rows
RATE_FIRST, RATE_LAST = 8, 17  # Settings: mileage-rate table
CAT_FIRST, CAT_LAST = 8, 21    # Settings: category list (12 + 2 custom)

CATEGORIES = [
    "Shipping Supplies / Materiales de Envío",
    "Postage & Shipping Labels / Envíos y Etiquetas",
    "Platform & Subscription Fees / Tarifas y Suscripciones",
    "Advertising & Promoted Listings / Publicidad",
    "Office Supplies / Artículos de Oficina",
    "Equipment & Tools / Equipo y Herramientas",
    "Phone & Internet / Teléfono e Internet",
    "Storage & Rent / Almacenaje y Renta",
    "Cleaning & Repairs / Limpieza y Reparaciones",
    "Software & Apps / Software y Aplicaciones",
    "Education & Courses / Educación y Cursos",
    "Other / Otro",
]
PAYMENTS = ["Card / Tarjeta", "Cash / Efectivo", "PayPal", "Bank / Banco", "Other / Otro"]
PURPOSES = ["Sourcing / Compras", "Post Office / Correo", "Supplies / Materiales", "Bank / Banco", "Other / Otro"]

wb = Workbook()
build_cover(
    wb, "Reseller Expense & Mileage Tracker", "Registro de Gastos y Millaje para Revendedores",
    [
        ("Dashboard", "📊  Dashboard / Panel", "Tax-year totals, deductions by category, monthly chart / Totales del año fiscal, deducciones por categoría, gráfica mensual"),
        ("Expenses", "🧾  Expenses / Gastos", "Log every business expense – deductible amount auto-calculated / Registra cada gasto – monto deducible automático"),
        ("Mileage", "🚗  Mileage / Millaje", "Log trips – miles × IRS rate by date, automatically / Registra viajes – millas × tarifa del IRS por fecha, automático"),
        ("Settings", "⚙️  Settings / Configuración", "IRS mileage rates & your expense categories / Tarifas del IRS y tus categorías de gastos"),
        ("Quick Start Guide", "🚀  Quick Start Guide / Guía Rápida", "Step-by-step setup in 5 minutes / Configuración paso a paso en 5 minutos"),
    ],
    "💻 PC / Mac  •  📱 iPhone / Android  —  Excel • Google Sheets   |   Computadora y celular — Excel • Google Sheets",
)

# =====================================================================
# SETTINGS
# =====================================================================
st = wb.create_sheet("Settings")
banner(st, "F", "Settings  /  Configuración", "Mileage rates and expense categories.  •  Tarifas por milla y categorías de gastos.")
for col, w in {"B": 16, "C": 16, "D": 36, "E": 3, "F": 50}.items():
    st.column_dimensions[col].width = w

section_title(st, "B6:D6", "IRS Mileage Rates  /  Tarifas por Milla del IRS")
st.row_dimensions[7].height = 36
header_cell(st["B7"], "Effective From\nVigente Desde")
header_cell(st["C7"], "Rate $ / mile\nTarifa $ / milla")
header_cell(st["D7"], "Source / Fuente")
rates = [
    (date(2025, 1, 1), 0.70, "IRS – 2025 business rate"),
    (date(2026, 1, 1), 0.725, "IRS Notice 2026-10"),
    (date(2026, 7, 1), 0.76, "IRS Announcement 2026-11 (mid-year increase)"),
]
for r in range(RATE_FIRST, RATE_LAST + 1):
    for col in "BCD":
        c = st[f"{col}{r}"]
        c.border = BORDER
        c.fill = fill("FEF9C3")
        c.font = font(10, color="1E3A8A")
    st[f"B{r}"].number_format = DATE_FMT
    st[f"B{r}"].alignment = Alignment(horizontal="center")
    st[f"C{r}"].number_format = "$0.000"
    st[f"C{r}"].alignment = Alignment(horizontal="center")
for i, (d, rate, src) in enumerate(rates):
    r = RATE_FIRST + i
    st[f"B{r}"], st[f"C{r}"], st[f"D{r}"] = d, rate, src
st.merge_cells(f"B{RATE_LAST + 1}:D{RATE_LAST + 3}")
st[f"B{RATE_LAST + 1}"] = ("Keep dates oldest → newest. When the IRS announces a new rate, add it on the next empty row – each trip uses the rate in effect on its date.  "
                           "Using kilometres? Enter your rate per km and log km in the Mileage sheet.\n"
                           "Mantén las fechas de la más antigua a la más reciente. Cuando el IRS anuncie una nueva tarifa, agrégala en la siguiente fila vacía. ¿Usas kilómetros? Escribe tu tarifa por km.")
st[f"B{RATE_LAST + 1}"].font = font(8, False, SLATE, italic=True)
st[f"B{RATE_LAST + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
st["C8"].comment = note("Rates as published by the IRS (verify at irs.gov before filing).\n"
                        "Tarifas publicadas por el IRS (verifica en irs.gov antes de declarar).")
date_validation(st, f"B{RATE_FIRST}:B{RATE_LAST}")
number_validation(st, [f"C{RATE_FIRST}:C{RATE_LAST}"], high="10", error="Enter the rate in dollars, e.g. 0.76.\nEscribe la tarifa en dólares, ej. 0.76.")

section_title(st, "F6:F6", "Expense Categories  /  Categorías de Gastos")
st.row_dimensions[7].height = 36
header_cell(st["F7"], "Category (edit or add your own)\nCategoría (edita o agrega)")
for i in range(CAT_LAST - CAT_FIRST + 1):
    r = CAT_FIRST + i
    c = st[f"F{r}"]
    c.value = CATEGORIES[i] if i < len(CATEGORIES) else None
    c.border = BORDER
    c.fill = fill("FEF9C3")
    c.font = font(10, color="1E3A8A")
st[f"F{CAT_LAST + 1}"] = "Rows 20–21 are free for your own categories. / Las filas 20–21 son para tus propias categorías."
st[f"F{CAT_LAST + 1}"].font = font(8, False, SLATE, italic=True)

# =====================================================================
# EXPENSES
# =====================================================================
ex = wb.create_sheet("Expenses", 1)
banner(ex, "L", "Expenses  /  Gastos",
       "Type in the light-teal columns only. Grey columns calculate automatically.  •  "
       "Escribe solo en las columnas verde claro. Las columnas grises se calculan solas.")
ex_headers = [
    ("B", "Date\nFecha", 13, "in"),
    ("C", "Vendor / Store\nProveedor / Tienda", 20, "in"),
    ("D", "Description\nDescripción", 26, "in"),
    ("E", "Category\nCategoría", 36, "in"),
    ("F", "Payment Method\nMétodo de Pago", 16, "in"),
    ("G", "Amount\nMonto", 13, "in"),
    ("H", "Business Use %\n% Uso del Negocio", 14, "in"),
    ("I", "Deductible $\nDeducible $", 14, "calc"),
    ("J", "Receipt Saved?\n¿Recibo Guardado?", 14, "in"),
    ("K", "Notes\nNotas", 22, "in"),
    ("L", "Check\nRevisar", 30, "calc"),
]
style_grid(ex, ex_headers, 5, FIRST, LAST, money="GI", pct="H", dates="B", centered="JL")
for r in range(FIRST, LAST + 1):
    ex[f"I{r}"] = f'=IF(G{r}="","",G{r}*IF(H{r}="",1,H{r}))'
    ex[f"L{r}"] = (f'=IF(G{r}="","",IF(ISNUMBER(B{r}),IF(E{r}="","⚠ ADD CATEGORY / AGREGA CATEGORÍA",'
                   f'IF(J{r}="{NO}","⚠ SAVE RECEIPT / GUARDA RECIBO","✔ OK")),"⚠ ADD DATE / AGREGA FECHA"))')
    ex[f"N{r}"] = f'=IF(G{r}="",0,IF(ISNUMBER(B{r}),1,0))'
    ex[f"O{r}"] = f'=IF(N{r}=1,YEAR(B{r}),0)'
    ex[f"P{r}"] = f'=IF(N{r}=1,YEAR(B{r})*100+MONTH(B{r}),0)'
    ex[f"Q{r}"] = f'=IF(N{r}=1,IF(J{r}="{NO}",1,0),0)'
hide_helpers(ex, 5, [("N", "valid (1/0)"), ("O", "year"), ("P", "year-month"), ("Q", "no receipt (1/0)")], FIRST, LAST)
ex["H5"].comment = note("Blank = 100%. For mixed personal/business items (e.g. your phone) enter the business share, e.g. 50%.\n"
                        "Vacío = 100%. Para artículos de uso mixto (ej. tu teléfono) escribe la parte del negocio, ej. 50%.")
ex["I5"].comment = note("Auto: Amount × Business Use %\nAuto: Monto × % Uso del Negocio")
ex["L5"].comment = note("Flags a missing date, category or receipt.\nMarca si falta la fecha, la categoría o el recibo.")

ex_samples = [
    (date(2026, 1, 5), "Uline", "Poly mailers 100-pack", 0, 0, 24.99, None, YES, ""),
    (date(2026, 1, 12), "eBay", "Store subscription – January", 2, 0, 27.95, None, YES, "Basic store"),
    (date(2026, 1, 20), "Staples", "Printer ink + labels", 4, 0, 38.40, None, YES, ""),
    (date(2026, 2, 2), "Verizon", "Phone bill – February", 6, 0, 80.00, 0.5, YES, "50% business"),
    (date(2026, 2, 14), "Amazon", "Digital shipping scale", 5, 0, 29.99, None, YES, ""),
    (date(2026, 3, 1), "Poshmark", "Closet promotion", 3, 2, 15.00, None, NO, ""),
    (date(2026, 3, 18), "USPS", "Postage – unreimbursed labels", 1, 1, 42.75, None, YES, ""),
    (date(2026, 4, 3), "U-Haul", "Storage unit – April", 7, 3, 65.00, None, YES, ""),
    (date(2026, 4, 22), "Home Depot", "Shelving + bins", 5, 0, 112.48, None, YES, "Inventory shelves"),
    (date(2026, 5, 9), "Vendoo", "Crosslisting app – monthly", 9, 0, 29.00, None, YES, ""),
    (date(2026, 5, 27), "Dollar Tree", "Tissue paper & tape", 0, 1, 11.25, None, NO, ""),
    (date(2026, 6, 15), "Udemy", "Reselling course", 10, 0, 19.99, None, YES, ""),
    (date(2026, 7, 8), "Walmart", "Steamer for clothes", 8, 0, 34.97, None, YES, ""),
    (date(2026, 7, 30), "Comcast", "Internet – July", 6, 0, 70.00, 0.3, YES, "30% business"),
    (date(2026, 8, 12), "Uline", "Boxes – assorted", 0, 0, 58.60, None, YES, ""),
    (date(2026, 8, 25), "Facebook", "Boosted Marketplace post", 3, 0, 10.00, None, YES, ""),
    (date(2026, 9, 10), "Canva", "Canva Pro – photo editing", 9, 0, 14.99, None, YES, ""),
    (date(2026, 9, 24), "Office Depot", "Label printer paper", 4, 0, 22.49, None, None, ""),
]
for i, (d, vendor, desc, cat, pay, amt, bus, rec, notes) in enumerate(ex_samples):
    r = FIRST + i
    ex[f"B{r}"], ex[f"C{r}"], ex[f"D{r}"] = d, vendor, desc
    ex[f"E{r}"], ex[f"F{r}"], ex[f"G{r}"] = CATEGORIES[cat], PAYMENTS[pay], amt
    ex[f"H{r}"], ex[f"J{r}"], ex[f"K{r}"] = bus, rec, notes or None
EX_SAMPLE_LAST = FIRST + len(ex_samples) - 1
ex["B6"].comment = note(f"Rows 6–{EX_SAMPLE_LAST} are EXAMPLES. Select B6:K{EX_SAMPLE_LAST} and press Delete to start fresh.\n"
                        f"Las filas 6–{EX_SAMPLE_LAST} son EJEMPLOS. Selecciona B6:K{EX_SAMPLE_LAST} y presiona Suprimir.")

list_validation(ex, f"E{FIRST}:E{LAST}", source=f"=Settings!$F${CAT_FIRST}:$F${CAT_LAST}", title="Category / Categoría")
list_validation(ex, f"F{FIRST}:F{LAST}", PAYMENTS, title="Payment / Pago")
list_validation(ex, f"J{FIRST}:J{LAST}", [YES, NO], title="Receipt / Recibo")
date_validation(ex, f"B{FIRST}:B{LAST}")
number_validation(ex, [f"G{FIRST}:G{LAST}"])
number_validation(ex, [f"H{FIRST}:H{LAST}"], high="1", title="Business % / % Negocio",
                  error="Enter a percent between 0% and 100%.\nEscribe un porcentaje entre 0% y 100%.")
contains_rule(ex, f"L{FIRST}:L{LAST}", "ADD", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(ex, f"L{FIRST}:L{LAST}", "SAVE", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(ex, f"L{FIRST}:L{LAST}", "OK", Font(name=FONT, color=GREEN), fill(GREEN_LIGHT))
ex.freeze_panes = "C6"
ex.auto_filter.ref = f"B5:L{LAST}"
ex.print_title_rows = "5:5"

# =====================================================================
# MILEAGE
# =====================================================================
mi = wb.create_sheet("Mileage", 2)
banner(mi, "N", "Mileage Log  /  Registro de Millaje",
       "Enter odometer readings OR the trip miles. The IRS rate for the trip date is applied automatically.  •  "
       "Escribe el odómetro O las millas del viaje. La tarifa del IRS de esa fecha se aplica sola.")
mi_headers = [
    ("B", "Date\nFecha", 13, "in"),
    ("C", "From\nDesde", 18, "in"),
    ("D", "To\nHasta", 18, "in"),
    ("E", "Purpose\nPropósito", 20, "in"),
    ("F", "Odometer Start\nOdómetro Inicial", 14, "in"),
    ("G", "Odometer End\nOdómetro Final", 14, "in"),
    ("H", "Miles (no odometer)\nMillas (sin odómetro)", 15, "in"),
    ("I", "Round Trip?\n¿Ida y Vuelta?", 12, "in"),
    ("J", "Total Miles\nMillas Totales", 12, "calc"),
    ("K", "Rate $/mile\nTarifa $/milla", 12, "calc"),
    ("L", "Deduction $\nDeducción $", 13, "calc"),
    ("M", "Check\nRevisar", 30, "calc"),
    ("N", "Notes\nNotas", 22, "in"),
]
style_grid(mi, mi_headers, 5, FIRST, LAST, money="L", dates="B", centered="IM", ints="FGHJ")
RB = f"Settings!$B${RATE_FIRST}:$B${RATE_LAST}"
RC = f"Settings!$C${RATE_FIRST}:$C${RATE_LAST}"
for r in range(FIRST, LAST + 1):
    mi[f"J{r}"] = (f'=IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r})),MAX(0,G{r}-F{r}),'
                   f'IF(ISNUMBER(H{r}),H{r}*IF(I{r}="{YES}",2,1),""))')
    n_rates = f'COUNTIF({RB},"<="&B{r})'
    mi[f"K{r}"] = (f'=IF(OR(J{r}="",NOT(ISNUMBER(B{r}))),"",'
                   f'IF({n_rates}=0,Settings!$C${RATE_FIRST},INDEX({RC},{n_rates})))')
    mi[f"K{r}"].number_format = "$0.000"
    mi[f"L{r}"] = f'=IF(K{r}="","",ROUND(J{r}*K{r},2))'
    mi[f"M{r}"] = (f'=IF(AND(B{r}="",F{r}="",G{r}="",H{r}=""),"",IF(NOT(ISNUMBER(B{r})),"⚠ ADD DATE / AGREGA FECHA",'
                   f'IF(J{r}="","⚠ ADD MILES / AGREGA MILLAS",IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r})),'
                   f'IF(G{r}<F{r},"⚠ CHECK ODOMETER / REVISA ODÓMETRO","✔ OK"),"✔ OK"))))')
    mi[f"P{r}"] = f'=IF(ISNUMBER(L{r}),1,0)'
    mi[f"Q{r}"] = f'=IF(P{r}=1,YEAR(B{r}),0)'
    mi[f"R{r}"] = f'=IF(P{r}=1,YEAR(B{r})*100+MONTH(B{r}),0)'
hide_helpers(mi, 5, [("P", "valid (1/0)"), ("Q", "year"), ("R", "year-month")], FIRST, LAST)
mi["H5"].comment = note("Use this when you don't record the odometer (e.g. miles from Google Maps).\n"
                        "Úsalo si no anotas el odómetro (ej. millas de Google Maps).")
mi["I5"].comment = note("'Yes' doubles the miles typed in 'Miles (no odometer)'. Odometer trips are never doubled.\n"
                        "'Sí' duplica las millas escritas en 'Millas (sin odómetro)'. Los viajes con odómetro no se duplican.")
mi["K5"].comment = note("Auto: the IRS rate in effect on the trip date (from Settings).\nAuto: la tarifa del IRS vigente en la fecha del viaje (de Configuración).")

mi_samples = [
    (date(2026, 1, 10), "Home", "Goodwill – Main St", 0, None, None, 6.2, YES),
    (date(2026, 1, 24), "Home", "Estate sale – Oak Ave", 0, 48210, 48236, None, None),
    (date(2026, 2, 7), "Home", "Post office", 1, None, None, 2.5, YES),
    (date(2026, 2, 21), "Home", "Flea market", 0, None, None, 14.0, YES),
    (date(2026, 3, 14), "Home", "Salvation Army + Savers", 0, 48702, 48741, None, None),
    (date(2026, 4, 4), "Home", "Staples", 2, None, None, 4.1, YES),
    (date(2026, 4, 25), "Home", "Garage sales route", 0, None, None, 22.5, None),
    (date(2026, 5, 16), "Home", "Post office", 1, None, None, 2.5, YES),
    (date(2026, 6, 6), "Home", "Bins outlet", 0, 49655, 49702, None, None),
    (date(2026, 6, 27), "Home", "Bank deposit", 3, None, None, 3.0, YES),
    (date(2026, 7, 11), "Home", "Goodwill – Main St", 0, None, None, 6.2, YES),
    (date(2026, 7, 25), "Home", "Estate sale – Pine Rd", 0, 50310, 50348, None, None),
    (date(2026, 8, 8), "Home", "Post office", 1, None, None, 2.5, YES),
    (date(2026, 8, 29), "Home", "Antique mall", 0, None, None, 18.4, YES),
    (date(2026, 9, 12), "Home", "Uline pickup", 2, None, None, 11.0, YES),
    (date(2026, 9, 26), "Home", "Flea market", 0, None, None, 14.0, YES),
]
for i, (d, frm, to, purpose, o1, o2, miles, rt) in enumerate(mi_samples):
    r = FIRST + i
    mi[f"B{r}"], mi[f"C{r}"], mi[f"D{r}"], mi[f"E{r}"] = d, frm, to, PURPOSES[purpose]
    mi[f"F{r}"], mi[f"G{r}"], mi[f"H{r}"], mi[f"I{r}"] = o1, o2, miles, rt
MI_SAMPLE_LAST = FIRST + len(mi_samples) - 1
mi["B6"].comment = note(f"Rows 6–{MI_SAMPLE_LAST} are EXAMPLES. Select B6:I{MI_SAMPLE_LAST} (and N) and press Delete.\n"
                        f"Las filas 6–{MI_SAMPLE_LAST} son EJEMPLOS. Selecciona B6:I{MI_SAMPLE_LAST} (y N) y presiona Suprimir.")
list_validation(mi, f"E{FIRST}:E{LAST}", PURPOSES, title="Purpose / Propósito")
list_validation(mi, f"I{FIRST}:I{LAST}", [YES, NO], title="Round trip / Ida y vuelta")
date_validation(mi, f"B{FIRST}:B{LAST}")
number_validation(mi, [f"F{FIRST}:H{LAST}"])
contains_rule(mi, f"M{FIRST}:M{LAST}", "⚠", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(mi, f"M{FIRST}:M{LAST}", "OK", Font(name=FONT, color=GREEN), fill(GREEN_LIGHT))
mi.freeze_panes = "C6"
mi.auto_filter.ref = f"B5:N{LAST}"
mi.print_title_rows = "5:5"

# =====================================================================
# DASHBOARD
# =====================================================================
db = wb.create_sheet("Dashboard", 1)
banner(db, "N", "Dashboard  /  Panel de Control",
       "Your tax-year deductions at a glance. Updates automatically.  •  Tus deducciones del año fiscal de un vistazo. Se actualiza solo.")
for col, w in {"B": 34, "C": 13, "D": 15, "E": 15, "F": 13, "G": 3,
               "H": 13, "I": 13, "J": 13, "K": 13, "L": 13, "M": 13, "N": 3}.items():
    db.column_dimensions[col].width = w


def xr(col):
    return f"Expenses!${col}${FIRST}:${col}${LAST}"


def mr(col):
    return f"Mileage!${col}${FIRST}:${col}${LAST}"


YEAR = "$C$6"
db["B6"] = "Tax Year / Año Fiscal:"
db["B6"].font = font(11, True, NAVY)
db["B6"].alignment = Alignment(horizontal="right", vertical="center")
setting_cell(db["C6"], f'=IF(COUNT({xr("B")},{mr("B")})=0,YEAR(TODAY()),YEAR(MAX({xr("B")},{mr("B")})))', "0")
db["C6"].comment = note("Shows the year of your latest entry. Type any year (e.g. 2027) to change it.\n"
                        "Muestra el año de tu último registro. Escribe otro año (ej. 2027) para cambiarlo.")
db.merge_cells("D6:F6")
db["D6"] = "← type a year to change / escribe un año para cambiar"
db["D6"].font = font(8, False, SLATE, italic=True)
db.row_dimensions[6].height = 22

kpi_cards(db, 8, [
    ("B", "B", "DEDUCTIBLE EXPENSES", "Gastos Deducibles", f"=ROUND(SUMIFS({xr('I')},{xr('O')},{YEAR}),2)", CUR0,
     f"=SUMIFS({xr('G')},{xr('O')},{YEAR})", '$#,##0" spent in total / gastado en total"'),
    ("C", "D", "BUSINESS MILES", "Millas de Negocio", f"=SUMIFS({mr('J')},{mr('Q')},{YEAR})", "#,##0",
     f"=COUNTIF({mr('Q')},{YEAR})", '0" trips / viajes"'),
    ("E", "F", "MILEAGE DEDUCTION", "Deducción por Millaje", f"=ROUND(SUMIFS({mr('L')},{mr('Q')},{YEAR}),2)", CUR0,
     "=IF(C10=0,0,E10/C10)", '$0.000" avg / mile · prom / milla"'),
    ("H", "J", "TOTAL DEDUCTIONS", "Deducciones Totales", "=B10+E10", CUR0,
     f"=SUMIFS({xr('Q')},{xr('O')},{YEAR})", '0" missing receipts / recibos faltantes"'),
])

# ---- By category ----
section_title(db, "B13:F13", "Deductible Expenses by Category  /  Gastos Deducibles por Categoría")
db.row_dimensions[14].height = 30
for col, text in zip("BCDEF", ["Category\nCategoría", "Entries\nRegistros", "Spent\nGastado", "Deductible\nDeducible", "% of Total\n% del Total"]):
    header_cell(db[f"{col}14"], text)
CAT_ROW0 = 15
n_cats = CAT_LAST - CAT_FIRST + 1
rows = []
for i in range(n_cats):
    r = CAT_ROW0 + i
    s = CAT_FIRST + i
    crit = f"{xr('O')},{YEAR},{xr('E')},$B{r}"
    rows.append([
        f'=IF(Settings!$F${s}="","",Settings!$F${s})',
        f'=IF($B{r}="",0,COUNTIFS({crit}))',
        f'=IF($B{r}="",0,ROUND(SUMIFS({xr("G")},{crit}),2))',
        f'=IF($B{r}="",0,ROUND(SUMIFS({xr("I")},{crit}),2))',
        f'=IF($E${CAT_ROW0 + n_cats + 1}=0,0,E{r}/$E${CAT_ROW0 + n_cats + 1})',
    ])
UNCAT = CAT_ROW0 + n_cats
TOT = UNCAT + 1
rows.append([
    "Uncategorized / Sin categoría",
    f"=C{TOT}-SUM(C{CAT_ROW0}:C{UNCAT - 1})",
    f"=ROUND(D{TOT}-SUM(D{CAT_ROW0}:D{UNCAT - 1}),2)",
    f"=ROUND(E{TOT}-SUM(E{CAT_ROW0}:E{UNCAT - 1}),2)",
    f"=IF($E${TOT}=0,0,E{UNCAT}/$E${TOT})",
])
fmts = {"C": "#,##0", "D": CUR_SIMPLE, "E": CUR_SIMPLE, "F": "0.0%"}
table_rows(db, CAT_ROW0, rows, "BCDEF", fmts, bold_first=False)
total_row(db, TOT, "BCDEF", ["TOTAL", f"=COUNTIF({xr('O')},{YEAR})", f"=ROUND(SUMIFS({xr('G')},{xr('O')},{YEAR}),2)",
                             f"=ROUND(SUMIFS({xr('I')},{xr('O')},{YEAR}),2)", f"=IF(E{TOT}=0,0,1)"], fmts)
db[f"B{TOT + 1}"] = "Categories come from the Settings sheet. / Las categorías vienen de la hoja Configuración."
db[f"B{TOT + 1}"].font = font(8, False, SLATE, italic=True)

# ---- Monthly ----
M_TITLE = TOT + 3
section_title(db, f"B{M_TITLE}:F{M_TITLE}", "Monthly Deductions  /  Deducciones Mensuales")
M_HEAD = M_TITLE + 1
db.row_dimensions[M_HEAD].height = 30
for col, text in zip("BCDEF", ["Month\nMes", "Expenses\nGastos", "Miles\nMillas", "Mileage $\nMillaje $", "Total\nTotal"]):
    header_cell(db[f"{col}{M_HEAD}"], text)
M0 = M_HEAD + 1
rows = []
for m, label in enumerate(MONTHS, start=1):
    r = M0 + m - 1
    key = f"{YEAR}*100+{m}"
    rows.append([label,
                 f"=ROUND(SUMIFS({xr('I')},{xr('P')},{key}),2)",
                 f"=SUMIFS({mr('J')},{mr('R')},{key})",
                 f"=ROUND(SUMIFS({mr('L')},{mr('R')},{key}),2)",
                 f"=C{r}+E{r}"])
mfmts = {"C": CUR_SIMPLE, "D": "#,##0.0", "E": CUR_SIMPLE, "F": CUR_SIMPLE}
table_rows(db, M0, rows, "BCDEF", mfmts)
M_TOT = M0 + 12
total_row(db, M_TOT, "BCDEF", ["YEAR / AÑO", "=B10", "=C10", "=E10", f"=C{M_TOT}+E{M_TOT}"], mfmts)

# ---- Charts ----
bar = bar_chart("Deductible by Category / Deducible por Categoría", horizontal=True)
bar.add_data(Reference(db, min_col=5, min_row=14, max_row=UNCAT), titles_from_data=True)
bar.set_categories(Reference(db, min_col=2, min_row=CAT_ROW0, max_row=UNCAT))
color_series(bar.series[0], TEAL, TEAL_DARK)
bar.dataLabels.numFmt = "$#,##0"
bar.x_axis.scaling.orientation = "maxMin"   # first category on top
bar.y_axis.crosses = "max"                  # keep the value axis at the bottom
bar.gapWidth = 40
bar.x_axis.txPr = axis_text(750)
bar.height, bar.width = 10.5, 17
db.add_chart(bar, "H13")

col = BarChart()
col.type = "col"
col.grouping = "stacked"
col.overlap = 100
col.title = small_title("Monthly Deductions / Deducciones Mensuales")
col.style = 10
col.y_axis.numFmt = "$#,##0"
col.y_axis.majorGridlines = None
col.x_axis.delete = False
col.y_axis.delete = False
col.add_data(Reference(db, min_col=3, min_row=M_HEAD, max_row=M0 + 11), titles_from_data=True)
col.add_data(Reference(db, min_col=5, min_row=M_HEAD, max_row=M0 + 11), titles_from_data=True)
col.set_categories(Reference(db, min_col=2, min_row=M0, max_row=M0 + 11))
col.series[0].tx = SeriesLabel(v="Expenses / Gastos")
col.series[1].tx = SeriesLabel(v="Mileage / Millaje")
color_series(col.series[0], TEAL, TEAL_DARK)
color_series(col.series[1], NAVY_2, NAVY)
col.legend.position = "t"
col.height, col.width = 8.5, 17
db.add_chart(col, f"H{M_TITLE}")

# =====================================================================
# QUICK START GUIDE
# =====================================================================
steps = [
    (f"Clear the example data: on Expenses select B6:K{EX_SAMPLE_LAST} and press Delete; on Mileage select B6:I{MI_SAMPLE_LAST} and press Delete. Never delete whole rows or the grey columns.",
     f"Borra los datos de ejemplo: en Expenses selecciona B6:K{EX_SAMPLE_LAST} y presiona Suprimir; en Mileage selecciona B6:I{MI_SAMPLE_LAST} y presiona Suprimir. Nunca borres filas completas ni las columnas grises."),
    ("Check Settings. The IRS mileage table is filled in (2026: 72.5¢/mile Jan 1 – Jun 30 and 76¢/mile from Jul 1). When the IRS announces a new rate, add it on the next empty row. Rename or add expense categories if you like.",
     "Revisa Settings (Configuración). La tabla de tarifas del IRS ya está llena (2026: 72.5¢/milla del 1 ene al 30 jun y 76¢/milla desde el 1 jul). Cuando el IRS anuncie una nueva tarifa, agrégala en la siguiente fila vacía. Puedes renombrar o agregar categorías."),
    ("Log every business expense on Expenses: date, vendor, description, category (dropdown), payment method and amount. For mixed-use items (like your phone) enter the Business Use % – blank means 100%.",
     "Registra cada gasto del negocio en Expenses: fecha, proveedor, descripción, categoría (lista), método de pago y monto. Para artículos de uso mixto (como tu teléfono) escribe el % de Uso del Negocio – vacío significa 100%."),
    ("Mark 'Receipt Saved?' Yes or No. The Check column flags anything missing: date, category or receipt.",
     "Marca '¿Recibo Guardado?' Sí o No. La columna Revisar marca lo que falte: fecha, categoría o recibo."),
    ("Log every business trip on Mileage: date, from, to and purpose. Enter the odometer start & end, OR just the miles (e.g. from Google Maps). 'Round Trip? = Yes' doubles the miles you typed.",
     "Registra cada viaje de negocio en Mileage: fecha, desde, hasta y propósito. Escribe el odómetro inicial y final, O solo las millas (ej. de Google Maps). '¿Ida y Vuelta? = Sí' duplica las millas escritas."),
    ("The correct IRS rate for each trip date is applied automatically, and Deduction $ = Total Miles × Rate.",
     "La tarifa correcta del IRS para la fecha de cada viaje se aplica sola, y Deducción $ = Millas Totales × Tarifa."),
    ("Open the Dashboard: deductible expenses, business miles, mileage deduction and total deductions for the tax year, plus category and monthly charts. It's locked to protect formulas – only the yellow Tax Year cell can be changed.",
     "Abre el Dashboard (Panel): gastos deducibles, millas de negocio, deducción por millaje y deducciones totales del año fiscal, más gráficas por categoría y por mes. Está bloqueado para proteger las fórmulas – solo se puede cambiar la celda amarilla del Año Fiscal."),
    ("At tax time, share the Dashboard totals (and your receipts) with your tax preparer.",
     "En temporada de impuestos, comparte los totales del Panel (y tus recibos) con tu preparador de impuestos."),
    ("Don't log inventory purchases here – those are your cost of goods. Track them per item in ProfitTrack Profit Tracker or ProfitTrack Inventory.",
     "No registres aquí la compra de mercancía – eso es tu costo de ventas. Regístrala por artículo en ProfitTrack Profit Tracker o ProfitTrack Inventory."),
    ("Google Sheets: upload to Google Drive → Open with Google Sheets → File → Save as Google Sheets. Tip: save a blank copy first and back up monthly.",
     "Google Sheets: súbelo a Google Drive → Abrir con Google Sheets → Archivo → Guardar como Hojas de cálculo de Google. Consejo: guarda una copia en blanco y respalda cada mes."),
]
faq = ("FAQ  /  Preguntas Frecuentes", [
    ("Q: I use kilometres.  A: Put your rate per km in Settings and type km in the Mileage sheet – the math is the same.",
     "P: Uso kilómetros.  R: Pon tu tarifa por km en Settings y escribe km en la hoja Mileage – el cálculo es igual."),
    ("Q: Which trips count?  A: Generally trips for business – sourcing, post office, supplies, bank. Regular commuting usually doesn't. Ask your tax professional about your situation.",
     "P: ¿Qué viajes cuentan?  R: En general los viajes de negocio – compras, correo, materiales, banco. El traslado habitual normalmente no. Consulta a tu profesional de impuestos."),
    ("Q: Why can't I type on the Dashboard?  A: It's locked so formulas can't be erased by accident. To unlock (no password): Excel → Review → Unprotect Sheet.",
     "P: ¿Por qué no puedo escribir en el Panel?  R: Está bloqueado para no borrar fórmulas por accidente. Para desbloquear (sin contraseña): Excel → Revisar → Desproteger hoja."),
    ("Q: How many rows?  A: 500 expenses and 500 trips per file. Start a fresh copy each tax year.",
     "P: ¿Cuántas filas?  R: 500 gastos y 500 viajes por archivo. Usa una copia nueva cada año fiscal."),
])
disclaimer = ("Important  /  Importante", [
    ("This template helps you organize records. It is not tax, legal or accounting advice. Deduction rules depend on your situation – consult a qualified tax professional. Mileage rates: IRS Notice 2026-10 and Announcement 2026-11 (verify at irs.gov).",
     "Esta plantilla te ayuda a organizar tus registros. No es asesoría fiscal, legal ni contable. Las reglas de deducción dependen de tu situación – consulta a un profesional de impuestos calificado. Tarifas: IRS Notice 2026-10 y Announcement 2026-11 (verifica en irs.gov)."),
])
build_guide(wb, steps, [faq, disclaimer])

finish(wb, db, OUT, "ProfitTrack Expenses — Reseller Expense & Mileage Tracker")
