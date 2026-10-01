"""Build ProfitTrack Pricing — Handmade Product Pricing Calculator for Etsy sellers (bilingual EN/ES).

Price math (one product):
    COST = materials + labor + packaging + overhead + shipping label
    S    = shipping charged to the buyer,  f = % fees (transaction + processing [+ Offsite Ads])
    F    = fixed fees per order (processing fixed + listing fee)
    profit(P) = P + S - f*(P + S) - F - COST
    price for target margin m (profit = m*P):  P = (COST + F - S*(1 - f)) / (1 - f - m)
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from common import (AMBER, AMBER_LIGHT, BORDER, CALC_FILL, CUR, CUR0, FONT, GREEN, GREEN_LIGHT, INPUT_BLUE,
                    INPUT_FILL, NAVY, NAVY_2, NO, PCT, RED, RED_LIGHT, SLATE, TEAL, TEAL_DARK, TEAL_LIGHT, WHITE,
                    YELLOW, YES, banner, build_cover, build_guide, contains_rule, fill, finish, font, header_cell,
                    hide_helpers, kpi_cards, list_validation, note, number_validation, section_title, setting_cell,
                    style_grid)

OUT = "product/ProfitTrack_Handmade_Pricing_Calculator.xlsx"
M_FIRST, M_LAST = 6, 205       # materials library rows
P_FIRST, P_LAST = 12, 211      # product list rows (200)

wb = Workbook()
build_cover(
    wb, "Handmade Pricing Calculator", "Calculadora de Precios para Productos Hechos a Mano",
    [
        ("Calculator", "🧮  Calculator / Calculadora", "Price one product in detail – materials, time, fees, profit / Calcula un producto en detalle"),
        ("Product List", "📋  Product List / Lista de Productos", "Price your whole shop at once & check every margin / Todo tu catálogo y sus márgenes"),
        ("Materials", "🧵  Materials / Materiales", "Your supply library – cost per unit calculated / Tu lista de materiales – costo por unidad automático"),
        ("Settings", "⚙️  Settings / Configuración", "Hourly rate, Etsy fees, rounding / Tarifa por hora, comisiones de Etsy, redondeo"),
        ("Quick Start Guide", "🚀  Quick Start Guide / Guía Rápida", "Step-by-step setup in 5 minutes / Configuración paso a paso en 5 minutos"),
    ],
    "💻 PC / Mac  •  📱 iPhone / Android  —  Excel • Google Sheets   |   Computadora y celular — Excel • Google Sheets",
    kicker="MAKER TOOLKIT  •  KIT PARA CREADORES",
    platforms="Etsy  •  Shopify  •  Craft Fairs  •  Markets  •  Ferias y Mercados",
)

# =====================================================================
# SETTINGS
# =====================================================================
st = wb.create_sheet("Settings")
banner(st, "E", "Settings  /  Configuración", "Your rate, Etsy fees and rounding.  •  Tu tarifa, comisiones de Etsy y redondeo.")
for col, w in {"B": 52, "C": 14, "D": 3, "E": 58}.items():
    st.column_dimensions[col].width = w
section_title(st, "B6:C6", "Your Numbers & Etsy Fees  /  Tus Números y Comisiones")
settings = [
    ("rate", "Your hourly rate ($ / hour) / Tu tarifa por hora", 20, "$#,##0.00", "Pay yourself! What is one hour of your time worth?"),
    ("overhead", "Overhead per item ($) / Gastos generales por artículo", 0.50, "$#,##0.00", "Electricity, tools wear, software… spread per item."),
    ("margin", "Default target profit margin / Margen de ganancia objetivo", 0.30, "0%", "Profit as a % of the price, after all costs and fees."),
    ("trans", "Etsy transaction fee / Comisión de transacción", 0.065, "0.0%", "Etsy US: 6.5% of item price + shipping charged."),
    ("procp", "Payment processing fee % / Procesamiento de pago %", 0.03, "0.0%", "Etsy Payments US: 3% + $0.25 per order."),
    ("procf", "Payment processing fixed ($) / Procesamiento fijo", 0.25, "$#,##0.00", ""),
    ("listing", "Listing fee ($) / Tarifa de publicación", 0.20, "$#,##0.00", "$0.20 per listing / renewal."),
    ("ads", "Offsite Ads fee / Anuncios externos", 0.15, "0%", "15% (12% if you sell $10k+/year) – only on sales from Offsite Ads."),
    ("ads_on", "Include Offsite Ads in prices? / ¿Incluir anuncios externos?", NO, None, "Choose Yes to price safely for Offsite Ads sales."),
    ("round", "Round prices to .99? / ¿Redondear precios a .99?", YES, None, "Yes: $23.41 → $23.99"),
]
S = {}
for i, (key, label, val, fmt, tip) in enumerate(settings):
    r = 7 + i
    S[key] = f"Settings!$C${r}"
    st[f"B{r}"] = label
    st[f"B{r}"].font = font(10, True, NAVY)
    st[f"B{r}"].border = BORDER
    setting_cell(st[f"C{r}"], val, fmt)
    st[f"E{r}"] = tip
    st[f"E{r}"].font = font(9, False, SLATE, italic=True)
list_validation(st, "C15:C16", [YES, NO], title="Yes / No")
number_validation(st, ["C7:C8", "C12:C13"])
number_validation(st, ["C9:C11", "C14"], high="0.9", title="Percent / Porcentaje",
                  error="Enter a percent between 0% and 90%.\nEscribe un porcentaje entre 0% y 90%.")
st.merge_cells("B18:E20")
st["B18"] = ("Etsy fees shown are for US sellers (2026). Fees change and differ by country – check Etsy's Fees & Payments Policy and update these cells.\n"
             "Las comisiones son para vendedores de EE. UU. (2026). Cambian y varían por país – revisa la política de Etsy y actualiza estas celdas.")
st["B18"].font = font(8, False, SLATE, italic=True)
st["B18"].alignment = Alignment(wrap_text=True, vertical="top")
FEE_RATE = f'({S["trans"]}+{S["procp"]}+IF({S["ads_on"]}="{YES}",{S["ads"]},0))'
FIXED = f'({S["procf"]}+{S["listing"]})'


def rounded(raw):
    """.99 rounding that never goes below the raw price."""
    return (f'IF({S["round"]}="{YES}",IF(ROUNDUP({raw},0)-0.01>={raw},ROUNDUP({raw},0)-0.01,ROUNDUP({raw},0)+0.99),'
            f'ROUNDUP({raw},2))')


# =====================================================================
# MATERIALS LIBRARY
# =====================================================================
ma = wb.create_sheet("Materials")
banner(ma, "H", "Materials Library  /  Lista de Materiales",
       "Enter what you paid for a pack and how many units it has – the cost per unit is calculated.  •  "
       "Escribe lo que pagaste por el paquete y cuántas unidades trae – el costo por unidad se calcula solo.")
m_headers = [
    ("B", "Material\nMaterial", 30, "in"),
    ("C", "Supplier\nProveedor", 18, "in"),
    ("D", "Pack Price\nPrecio del Paquete", 14, "in"),
    ("E", "Units in Pack\nUnidades", 12, "in"),
    ("F", "Unit\nUnidad", 12, "in"),
    ("G", "Cost per Unit\nCosto por Unidad", 14, "calc"),
    ("H", "Notes\nNotas", 26, "in"),
]
style_grid(ma, m_headers, 5, M_FIRST, M_LAST, money="DG", centered="F")
for r in range(M_FIRST, M_LAST + 1):
    ma[f"E{r}"].number_format = "#,##0.##"
    ma[f"G{r}"] = f'=IF(OR(D{r}="",E{r}="",E{r}=0),"",D{r}/E{r})'
    ma[f"G{r}"].number_format = '$#,##0.000'
materials = [
    ("Soy wax", "CandleScience", 45.00, 160, "oz"),
    ("Fragrance oil", "CandleScience", 17.60, 16, "oz"),
    ("Cotton wick", "Amazon", 11.99, 100, "each"),
    ("Amber jar 8 oz", "Uline", 34.80, 24, "each"),
    ("Product label", "Sticker Mule", 30.00, 200, "each"),
    ("Sterling silver ear wires", "Rio Grande", 24.00, 50, "pair"),
    ("Glass beads", "Fire Mountain", 12.50, 250, "each"),
    ("Merino yarn", "LoveCrafts", 9.50, 1, "skein"),
    ("Vinyl sticker paper", "Cricut", 14.99, 20, "sheet"),
    ("Kraft gift box", "Uline", 42.00, 100, "each"),
    ("Tissue paper", "Dollar Tree", 1.25, 20, "sheet"),
    ("Thank-you card", "Vistaprint", 25.00, 250, "each"),
]
for i, (name, sup, price, qty, unit) in enumerate(materials):
    r = M_FIRST + i
    ma[f"B{r}"], ma[f"C{r}"], ma[f"D{r}"], ma[f"E{r}"], ma[f"F{r}"] = name, sup, price, qty, unit
UNITS = ["each", "pair", "oz", "lb", "g", "kg", "ml", "l", "inch", "ft", "yard", "m", "sheet", "skein", "hour"]
list_validation(ma, f"F{M_FIRST}:F{M_LAST}", UNITS, title="Unit / Unidad")
number_validation(ma, [f"D{M_FIRST}:E{M_LAST}"])
ma["B6"].comment = note("Rows 6–17 are EXAMPLES – replace them with your own supplies.\nLas filas 6–17 son EJEMPLOS – reemplázalas con tus materiales.")
ma.freeze_panes = "C6"

# =====================================================================
# CALCULATOR (one product, detailed)
# =====================================================================
ca = wb.create_sheet("Calculator", 1)
banner(ca, "J", "Pricing Calculator  /  Calculadora de Precios",
       "Price one product in detail. Type in the light-teal cells – everything else is calculated.  •  "
       "Calcula un producto en detalle. Escribe en las celdas verde claro – lo demás se calcula solo.")
for col, w in {"B": 34, "C": 12, "D": 10, "E": 13, "F": 14, "G": 13, "H": 3, "I": 36, "J": 18}.items():
    ca.column_dimensions[col].width = w


def label(cell, text, bold=True):
    ca[cell] = text
    ca[cell].font = font(10, bold, NAVY)
    ca[cell].border = BORDER
    ca[cell].alignment = Alignment(vertical="center", wrap_text=True)


def inp(cell, value=None, fmt=None):
    c = ca[cell]
    c.value = value
    c.fill = fill(INPUT_FILL)
    c.font = font(11, True, INPUT_BLUE)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt


def calc(cell, formula, fmt=None, bold=False):
    c = ca[cell]
    c.value = formula
    c.fill = fill(CALC_FILL)
    c.font = font(11, bold)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt


section_title(ca, "B6:G6", "1 · Product  /  Producto")
label("B7", "Product name / Nombre del producto")
ca.merge_cells("C7:G7")
inp("C7", "Soy Candle – 8 oz Amber Jar")

section_title(ca, "B9:G9", "2 · Materials  /  Materiales")
ca.row_dimensions[10].height = 36
for col, text in zip("BCDEFG", ["Material (pick from list)\nMaterial (elige de la lista)", "Qty Used\nCantidad", "Unit\nUnidad",
                                "Library Cost\nCosto de Lista", "Your Cost (optional)\nTu Costo (opcional)", "Line Cost\nCosto"]):
    header_cell(ca[f"{col}10"], text)
MAT_ROWS = range(11, 21)
LIB_B, LIB_G, LIB_F = (f"Materials!${c}${M_FIRST}:${c}${M_LAST}" for c in "BGF")
for r in MAT_ROWS:
    inp(f"B{r}")
    ca[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center")
    inp(f"C{r}", fmt="#,##0.##")
    calc(f"D{r}", f'=IF(B{r}="","",IFERROR(INDEX({LIB_F},MATCH(B{r},{LIB_B},0)),""))')
    calc(f"E{r}", f'=IF(B{r}="","",IFERROR(INDEX({LIB_G},MATCH(B{r},{LIB_B},0)),""))', "$#,##0.000")
    inp(f"F{r}", fmt="$#,##0.000")
    calc(f"G{r}", f'=IF(OR(B{r}="",C{r}=""),"",IF(F{r}<>"",C{r}*F{r},IF(E{r}="","",C{r}*E{r})))', CUR)
    ca[f"I{r}"] = f'=IF(AND(B{r}<>"",E{r}="",F{r}=""),"⚠ Not in Materials – type Your Cost / Escribe tu costo","")'
    ca[f"I{r}"].font = font(8, True, RED)
for r, (mat, qty) in zip(MAT_ROWS, [("Soy wax", 7), ("Fragrance oil", 0.6), ("Cotton wick", 1), ("Amber jar 8 oz", 1), ("Product label", 1)]):
    ca[f"B{r}"], ca[f"C{r}"] = mat, qty
list_validation(ca, "B11:B20", source=f"=Materials!$B${M_FIRST}:$B${M_LAST}", title="Material")
number_validation(ca, ["C11:C20", "F11:F20"])
label("B21", "MATERIALS TOTAL / TOTAL MATERIALES")
ca.merge_cells("C21:F21")
calc("G21", "=SUM(G11:G20)", CUR, bold=True)
ca["B10"].comment = note("Pick a material from your Materials sheet. Not in the list? Type its name and fill 'Your Cost'.\n"
                         "Elige un material de tu hoja Materiales. ¿No está? Escribe el nombre y llena 'Tu Costo'.")

section_title(ca, "B23:G23", "3 · Time & Other Costs  /  Tiempo y Otros Costos")
rows3 = [
    (24, "Minutes to make one / Minutos para hacer uno", 12, "0", "in"),
    (25, "Hourly rate / Tarifa por hora", f"={S['rate']}", "$#,##0.00", "in"),
    (26, "Labor cost / Costo de mano de obra", "=IF(C24=\"\",0,C24/60*C25)", CUR, "calc"),
    (27, "Packaging / Empaque", 0.80, CUR, "in"),
    (28, "Overhead per item / Gastos generales", f"={S['overhead']}", CUR, "in"),
    (29, "Shipping label you pay / Etiqueta de envío que pagas", 5.50, CUR, "in"),
    (30, "Shipping you charge the buyer (0 = free) / Envío que cobras", 0, CUR, "in"),
]
for r, text, val, fmt, kind in rows3:
    label(f"B{r}", text, bold=False)
    ca.merge_cells(f"C{r}:D{r}")
    (inp if kind == "in" else calc)(f"C{r}", val, fmt)
ca["C25"].comment = note("Filled from Settings – type a different rate for this product if you like.\nViene de Configuración – puedes escribir otra tarifa.")
ca["C30"].comment = note("Offering free shipping? Leave 0 – the label cost is then built into your price.\n"
                         "¿Envío gratis? Deja 0 – el costo de la etiqueta se incluye en el precio.")
number_validation(ca, ["C24", "C25", "C27:C30"])
label("B31", "TOTAL COST / COSTO TOTAL")
ca.merge_cells("C31:D31")
calc("C31", "=G21+C26+C27+C28+C29", CUR, bold=True)
ca["C31"].fill = fill(TEAL_LIGHT)

section_title(ca, "B33:G33", "4 · Your Target  /  Tu Objetivo")
label("B34", "Target profit margin / Margen objetivo", bold=False)
ca.merge_cells("C34:D34")
inp("C34", f"={S['margin']}", "0%")
label("B35", "Etsy % fees (from Settings) / Comisiones % de Etsy", bold=False)
ca.merge_cells("C35:D35")
calc("C35", f"={FEE_RATE}", "0.0%")
label("B36", "Etsy fixed fees per order / Comisiones fijas por pedido", bold=False)
ca.merge_cells("C36:D36")
calc("C36", f"={FIXED}", CUR)
number_validation(ca, ["C34"], high="0.85", title="Margin / Margen", error="Enter 0%–85%.\nEscribe 0%–85%.")

COST, SHIP, FR, FX, MG = "$C$31", "$C$30", "$C$35", "$C$36", "$C$34"


def fees(p):
    return f"(({p}+{SHIP})*{FR}+{FX})"


def profit(p):
    return f"({p}+{SHIP}-{fees(p)}-{COST})"


# ---- Results (right side) ----
section_title(ca, "I6:J6", "Your Prices  /  Tus Precios")
RAW = f"({COST}+{FX}-{SHIP}*(1-{FR}))/(1-{FR}-{MG})"
results = [
    (7, "BREAK-EVEN PRICE\nPrecio mínimo (sin ganancia)", f"=MAX(0,({COST}+{FX}-{SHIP}*(1-{FR}))/(1-{FR}))", NAVY_2, "Below this price you lose money.\nDebajo de este precio pierdes dinero."),
    (9, "RECOMMENDED PRICE\nPrecio recomendado", f"=IF(1-{FR}-{MG}<=0,0,MAX(0,{rounded(RAW)}))", TEAL_DARK, "Covers all costs, Etsy fees and your target margin.\nCubre costos, comisiones y tu margen objetivo."),
]
for r, text, f, color, tip in results:
    ca[f"I{r}"] = text
    ca[f"I{r}"].font = font(11, True, WHITE)
    ca[f"I{r}"].fill = fill(color)
    ca[f"I{r}"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ca[f"J{r}"] = f
    ca[f"J{r}"].font = font(20, True, WHITE)
    ca[f"J{r}"].fill = fill(color)
    ca[f"J{r}"].number_format = "$#,##0.00"
    ca[f"J{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ca[f"I{r}"].comment = note(tip)
    ca.row_dimensions[r].height = 44
REC = "$J$9"
detail = [
    (11, "Etsy fees at this price / Comisiones de Etsy", f"=IF({REC}=0,0,{fees(REC)})", CUR),
    (12, "Your profit per sale / Tu ganancia por venta", f"=IF({REC}=0,0,{profit(REC)})", CUR),
    (13, "Profit margin / Margen de ganancia", f"=IF({REC}=0,0,J12/{REC})", PCT),
    (14, "You earn per hour (profit + labor) / Ganas por hora", '=IF(OR(C24="",C24=0),0,(J12+C26)/(C24/60))', CUR),
]
for r, text, f, fmt in detail:
    ca[f"I{r}"] = text
    ca[f"I{r}"].font = font(10, False, NAVY)
    ca[f"I{r}"].border = BORDER
    calc(f"J{r}", f, fmt, bold=True)
ca["I15"] = f'=IF(1-{FR}-{MG}<=0,"⚠ Target margin too high for these fees / Margen objetivo demasiado alto","")'
ca["I15"].font = font(9, True, RED)

section_title(ca, "I17:J17", "Test Your Own Price  /  Prueba tu Precio")
ca["I18"] = "Your price / Tu precio"
ca["I18"].font = font(11, True, NAVY)
ca["I18"].border = BORDER
inp("J18", 29.99, "$#,##0.00")
number_validation(ca, ["J18"])
YP = "$J$18"
test = [
    (19, "Etsy fees / Comisiones", f'=IF({YP}="",0,{fees(YP)})', CUR),
    (20, "Profit per sale / Ganancia por venta", f'=IF({YP}="",0,{profit(YP)})', CUR),
    (21, "Profit margin / Margen", f'=IF(OR({YP}="",{YP}=0),0,J20/{YP})', PCT),
    (22, "You earn per hour / Ganas por hora", '=IF(OR(C24="",C24=0,J18=""),0,(J20+C26)/(C24/60))', CUR),
]
for r, text, f, fmt in test:
    ca[f"I{r}"] = text
    ca[f"I{r}"].font = font(10, False, NAVY)
    ca[f"I{r}"].border = BORDER
    calc(f"J{r}", f, fmt, bold=True)
ca.merge_cells("I23:J23")
ca["I23"] = (f'=IF({YP}="","",IF(J20<0,"⚠ LOSING MONEY / PIERDES DINERO",IF(J21<{MG},"● BELOW YOUR TARGET / DEBAJO DE TU OBJETIVO",'
             f'"✔ ON TARGET / EN TU OBJETIVO")))')
ca["I23"].font = font(12, True, NAVY)
ca["I23"].alignment = Alignment(horizontal="center", vertical="center")
ca.row_dimensions[23].height = 28
contains_rule(ca, "I23", "LOSING", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(ca, "I23", "BELOW", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(ca, "I23", "ON TARGET", Font(name=FONT, bold=True, color=GREEN), fill(GREEN_LIGHT))

section_title(ca, "I25:J25", "Classic Formula  /  Fórmula Clásica")
classic = [
    (26, "Wholesale = cost × 2 / Mayoreo = costo × 2", f"={COST}*2", CUR),
    (27, "Retail = wholesale × 2 / Menudeo = mayoreo × 2", f"={COST}*4", CUR),
    (28, "Profit at retail after Etsy fees / Ganancia al menudeo", f"={profit('J27')}", CUR),
]
for r, text, f, fmt in classic:
    ca[f"I{r}"] = text
    ca[f"I{r}"].font = font(10, False, NAVY)
    ca[f"I{r}"].border = BORDER
    calc(f"J{r}", f, fmt)
ca.merge_cells("I29:J31")
ca["I29"] = ("Compare: the classic 2×2 formula ignores Etsy fees and your time. The recommended price above includes both.\n"
             "Compara: la fórmula clásica ignora las comisiones y tu tiempo. El precio recomendado incluye ambos.")
ca["I29"].font = font(8, False, SLATE, italic=True)
ca["I29"].alignment = Alignment(wrap_text=True, vertical="top")
ca.freeze_panes = "A6"

# =====================================================================
# PRODUCT LIST (whole shop)
# =====================================================================
pl = wb.create_sheet("Product List", 2)
banner(pl, "R", "Product List  /  Lista de Productos",
       "Price your whole shop. Uses your hourly rate and Etsy fees from Settings.  •  "
       "Calcula todo tu catálogo con tu tarifa y las comisiones de Configuración.")
pheaders = [
    ("B", "Product\nProducto", 28, "in"),
    ("C", "Materials $\nMateriales $", 12, "in"),
    ("D", "Minutes\nMinutos", 10, "in"),
    ("E", "Packaging $\nEmpaque $", 11, "in"),
    ("F", "Ship Label $\nEtiqueta $", 11, "in"),
    ("G", "Ship Charged $\nEnvío Cobrado $", 12, "in"),
    ("H", "Target Margin\nMargen Objetivo", 12, "in"),
    ("I", "Total Cost\nCosto Total", 12, "calc"),
    ("J", "Break-even\nPrecio Mínimo", 12, "calc"),
    ("K", "Recommended\nRecomendado", 13, "calc"),
    ("L", "Profit @ Rec.\nGanancia", 12, "calc"),
    ("M", "Your Price\nTu Precio", 12, "in"),
    ("N", "Etsy Fees\nComisiones", 11, "calc"),
    ("O", "Your Profit\nTu Ganancia", 12, "calc"),
    ("P", "Your Margin\nTu Margen", 11, "calc"),
    ("Q", "$ / Hour\n$ / Hora", 11, "calc"),
    ("R", "Status\nEstado", 30, "calc"),
]
style_grid(pl, pheaders, 11, P_FIRST, P_LAST, money="CEFGIJKLMNOQ", pct="HP", centered="DR")
for r in range(P_FIRST, P_LAST + 1):
    pl[f"P{r}"].number_format = PCT
    m = f'IF(H{r}="",{S["margin"]},H{r})'
    cost = f"$I{r}"
    sh = f"$G{r}"

    def pfees(p):
        return f"(({p}+{sh})*{FEE_RATE}+{FIXED})"

    def pprofit(p):
        return f"({p}+{sh}-{pfees(p)}-{cost})"

    raw = f"({cost}+{FIXED}-{sh}*(1-{FEE_RATE}))/(1-{FEE_RATE}-{m})"
    pl[f"I{r}"] = f'=IF(B{r}="","",C{r}+IF(D{r}="",0,D{r}/60*{S["rate"]})+E{r}+F{r}+{S["overhead"]})'
    pl[f"J{r}"] = f'=IF(B{r}="","",MAX(0,({cost}+{FIXED}-{sh}*(1-{FEE_RATE}))/(1-{FEE_RATE})))'
    pl[f"K{r}"] = f'=IF(B{r}="","",IF(1-{FEE_RATE}-{m}<=0,0,MAX(0,{rounded(raw)})))'
    pl[f"L{r}"] = f'=IF(OR(B{r}="",K{r}=0),"",{pprofit(f"K{r}")})'
    pl[f"N{r}"] = f'=IF(OR(B{r}="",M{r}=""),"",{pfees(f"M{r}")})'
    pl[f"O{r}"] = f'=IF(OR(B{r}="",M{r}=""),"",{pprofit(f"M{r}")})'
    pl[f"P{r}"] = f'=IF(O{r}="","",IF(M{r}=0,0,O{r}/M{r}))'
    pl[f"Q{r}"] = f'=IF(OR(O{r}="",D{r}="",D{r}=0),"",(O{r}+D{r}/60*{S["rate"]})/(D{r}/60))'
    pl[f"R{r}"] = (f'=IF(B{r}="","",IF(M{r}="","Enter your price / Escribe tu precio",IF(O{r}<0,"⚠ LOSING MONEY / PIERDES",'
                   f'IF(P{r}<{m},"● BELOW TARGET / DEBAJO","✔ ON TARGET / EN OBJETIVO"))))')
    # helpers: priced (1/0), margin if priced, below target, losing
    pl[f"T{r}"] = f'=IF(B{r}="",0,IF(M{r}="",0,1))'
    pl[f"U{r}"] = f'=IF(T{r}=1,P{r},0)'
    pl[f"V{r}"] = f'=IF(T{r}=1,IF(O{r}<0,0,IF(P{r}<{m},1,0)),0)'
    pl[f"W{r}"] = f'=IF(T{r}=1,IF(O{r}<0,1,0),0)'
    pl[f"X{r}"] = f'=IF(T{r}=1,M{r},0)'
    pl[f"Y{r}"] = f'=IF(T{r}=1,O{r},0)'
    pl[f"Z{r}"] = f'=IF(B{r}="",0,1)'
hide_helpers(pl, 11, [("T", "priced"), ("U", "margin"), ("V", "below target"), ("W", "losing"),
                      ("X", "price"), ("Y", "profit"), ("Z", "listed")], P_FIRST, P_LAST)
products = [
    ("Soy Candle 8 oz", 4.35, 12, 0.80, 5.50, 0, None, 29.99),
    ("Beaded Drop Earrings", 2.10, 25, 0.60, 4.20, 0, None, 28.00),
    ("Merino Crochet Beanie", 9.50, 150, 1.00, 5.80, 0, None, 45.00),
    ("Vinyl Sticker Pack (5)", 1.20, 10, 0.30, 1.00, 0, None, 12.00),
    ("Personalized Wood Sign", 6.75, 60, 2.50, 9.50, 0, 0.35, 49.00),
    ("Goat Milk Soap Bar", 1.35, 8, 0.40, 4.50, 4.50, None, 10.50),
    ("Resin Keychain", 1.80, 20, 0.30, 4.20, 0, None, 12.00),
    ("Embroidered Tote Bag", 4.20, 45, 0.80, 6.10, 0, None, 42.00),
]
for i, prod in enumerate(products):
    r = P_FIRST + i
    for col, v in zip("BCDEFGHM", prod):
        pl[f"{col}{r}"] = v
pl["H11"].comment = note("Blank = your default target margin from Settings.\nVacío = tu margen objetivo de Configuración.")
pl["B12"].comment = note("Rows 12–19 are EXAMPLES – replace them with your products.\nLas filas 12–19 son EJEMPLOS – reemplázalas con tus productos.")
number_validation(pl, [f"C{P_FIRST}:G{P_LAST}", f"M{P_FIRST}:M{P_LAST}"])
number_validation(pl, [f"H{P_FIRST}:H{P_LAST}"], high="0.85", title="Margin / Margen", error="Enter 0%–85%.\nEscribe 0%–85%.")
contains_rule(pl, f"R{P_FIRST}:R{P_LAST}", "LOSING", Font(name=FONT, bold=True, color=RED), fill(RED_LIGHT))
contains_rule(pl, f"R{P_FIRST}:R{P_LAST}", "BELOW", Font(name=FONT, bold=True, color=AMBER), fill(AMBER_LIGHT))
contains_rule(pl, f"R{P_FIRST}:R{P_LAST}", "ON TARGET", Font(name=FONT, color=GREEN), fill(GREEN_LIGHT))


def pr(col):
    return f"${col}${P_FIRST}:${col}${P_LAST}"


kpi_cards(pl, 6, [
    ("B", "C", "PRODUCTS PRICED", "Productos con Precio", f"=SUM({pr('T')})", "#,##0", f"=SUM({pr('Z')})", '0" products listed / productos en la lista"'),
    ("E", "G", "AVERAGE MARGIN", "Margen Promedio", f"=IF(SUM({pr('X')})=0,0,SUM({pr('Y')})/SUM({pr('X')}))", PCT,
     "weighted by price / ponderado", None),
    ("I", "K", "BELOW TARGET", "Debajo del Objetivo", f"=SUM({pr('V')})", "#,##0", "raise these prices / sube estos precios", None),
    ("M", "O", "LOSING MONEY", "Perdiendo Dinero", f"=SUM({pr('W')})", "#,##0", "fix these first / arregla estos primero", None),
])
pl.freeze_panes = "C12"
pl.auto_filter.ref = f"B11:R{P_LAST}"
pl.print_title_rows = "11:11"

# =====================================================================
# QUICK START GUIDE
# =====================================================================
steps = [
    ("Open Settings: enter your hourly rate and overhead per item, choose your default target margin, and check the Etsy fees (US 2026 rates are filled in).",
     "Abre Settings (Configuración): escribe tu tarifa por hora y gastos generales, elige tu margen objetivo y revisa las comisiones de Etsy (tarifas de EE. UU. 2026 ya incluidas)."),
    ("Fill the Materials sheet with your supplies: what you paid for a pack and how many units it has. The cost per unit is calculated for you. Replace the example rows.",
     "Llena la hoja Materials con tus materiales: lo que pagaste por el paquete y cuántas unidades trae. El costo por unidad se calcula solo. Reemplaza los ejemplos."),
    ("Calculator – price one product in detail: pick each material from the list and enter the quantity used. Not in your list? Type the name and fill 'Your Cost'.",
     "Calculator – calcula un producto en detalle: elige cada material de la lista y escribe la cantidad usada. ¿No está? Escribe el nombre y llena 'Tu Costo'."),
    ("Add your time (minutes), packaging, shipping label cost and what you charge for shipping (0 = free shipping, built into the price).",
     "Agrega tu tiempo (minutos), empaque, costo de la etiqueta y lo que cobras de envío (0 = envío gratis, incluido en el precio)."),
    ("Read your prices: Break-even (no profit), Recommended (covers costs, Etsy fees and your target margin, rounded to .99), your profit per sale and what you earn per hour.",
     "Lee tus precios: Mínimo (sin ganancia), Recomendado (cubre costos, comisiones y tu margen, redondeado a .99), tu ganancia por venta y cuánto ganas por hora."),
    ("Test Your Own Price: type any price to see the fees, profit, margin and a ✔ / ● / ⚠ verdict instantly.",
     "Prueba tu Precio: escribe cualquier precio y verás comisiones, ganancia, margen y un veredicto ✔ / ● / ⚠ al instante."),
    ("Product List – price your whole shop: one row per product with materials $, minutes, packaging and shipping. Enter your current price to see your real margin and $ per hour.",
     "Product List – todo tu catálogo: una fila por producto con materiales $, minutos, empaque y envío. Escribe tu precio actual para ver tu margen real y $ por hora."),
    ("Fix the red ⚠ LOSING MONEY products first, then raise the ● BELOW TARGET ones. The summary cards at the top count them for you.",
     "Arregla primero los productos ⚠ PIERDES DINERO y luego sube los ● DEBAJO DEL OBJETIVO. Las tarjetas de arriba los cuentan por ti."),
    ("Selling through Offsite Ads? Set 'Include Offsite Ads' to Yes in Settings to price safely for those sales.",
     "¿Vendes con Anuncios Externos? Cambia 'Incluir anuncios externos' a Sí en Configuración para cubrir esas ventas."),
    ("Google Sheets: upload to Google Drive → Open with Google Sheets → File → Save as Google Sheets. Tip: save a blank copy first.",
     "Google Sheets: súbelo a Google Drive → Abrir con Google Sheets → Archivo → Guardar como Hojas de cálculo de Google. Consejo: guarda una copia en blanco."),
]
faq = ("FAQ  /  Preguntas Frecuentes", [
    ("Q: What's the difference between markup and margin?  A: Margin is profit as a % of the price. A 30% margin on a $30 item means $9 profit after all costs and fees.",
     "P: ¿Diferencia entre markup y margen?  R: El margen es la ganancia como % del precio. Un margen de 30% en un artículo de $30 son $9 de ganancia después de costos y comisiones."),
    ("Q: Do I include shipping?  A: Yes – the label you pay is a cost. If you charge the buyer for shipping, enter it too: Etsy charges fees on it, and the calculator handles that.",
     "P: ¿Incluyo el envío?  R: Sí – la etiqueta que pagas es un costo. Si cobras envío, escríbelo también: Etsy cobra comisión sobre él y la calculadora lo considera."),
    ("Q: I sell outside the US.  A: Change the fee cells in Settings to your country's Etsy rates. Prices are shown with a $ sign but work for any currency.",
     "P: Vendo fuera de EE. UU.  R: Cambia las comisiones en Settings a las de tu país. Los precios muestran $ pero funcionan con cualquier moneda."),
])
build_guide(wb, steps, [faq])

finish(wb, None, OUT, "ProfitTrack Pricing — Handmade Product Pricing Calculator")
